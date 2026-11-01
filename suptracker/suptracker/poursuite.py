"""Sprint 6: success beyond licence. BUT (per university) and master (per university x discipline).

Sources: five SIES workbooks the owner downloads from the ministry site (it rejects automated
downloads) and commits under data/sources/. They are recognised by their CONTENT (sheet names and
table titles), not by their file names, so a renamed or next-year file still loads. What SIES
publishes there, checked on the real files (NF 2025-26 for BUT, NF 2025-30 for master):

BUT ("Tableaux par établissement", "Tableaux réussite en IUT")
* success in 3 years of the 2021 bacheliers entering BUT 1, PER UNIVERSITY only: no IUT, no
  specialty. Success is credited to the university of first enrolment, wherever the diploma was
  obtained. Observed rate, "simulated" rate (expected from the entrants' profile, the specialty of
  BUT included) and value added = observed - simulated;
* nothing is published below 100 entrants ("ns");
* the continuation of the graduates in 3 years, in 2024, per university: share continuing, and
  the shares (of graduates) in master, business school, engineering school;
* national rates only by specialty (24), by bac series and général doublette, by bac mention.
  These are margins: never crossed (no "MATHS+SES with Assez bien" figure exists).

Master ("Réussite et passage par discipline et établissement", "Réussite en master",
"Poursuites d'études en master")
* per university x broad discipline of M1 (the five B1 disciplines): success in 2 or 3 years of
  the 2021 M1 entrants; M1 -> M2 transition of the 2023 entrants with repeat, reorientations and
  exits, a partition of the entrants. The population is EVERY student entering M1 at that
  university, not the licence graduates of this licence;
* national: continuation into master of licence graduates by discipline (2021-2024), and the full
  chain of the 2017 L1 entrants (licence in 3-4 years, master in 5-7 years) by bac type and by
  mention, again margins only.

As in B1, discipline rows may carry a shorter university name than their block's total row, and
the same university is spelled differently from one table to the next: rows are grouped by block,
and every spelling maps to the B1 university name through NAME_ALIASES (red team P2 checks that
no two blocks of a table collide). Universities add up to the "Inscriptions principales" row,
which is the national reference ("Etudiants" counts a person once and does not add up).

Official rates are stored and shown as published, never recomputed.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

from . import config
from . import evidence as E
from . import success as SU
from .success import fold

NATIONAL = SU.NATIONAL
SMALL_COHORT = 100
SOURCE_LABEL = {"but": "SIES, Note Flash 2025-26", "master": "SIES, Note Flash 2025-30"}

# Spellings of the new workbooks -> the B1 university name (keys of success.SIES_PAYSAGE).
NAME_ALIASES = {
    "CY CERGY PARIS UNIVERSITE": "CY CERGY UNIVERSITE",
    "CY CERGY PARIS UNIVERSITE (EPE)": "CY CERGY UNIVERSITE",
    "UNIV GRENOBLE ALPES (EPE)": "UNIVERSITE GRENOBLE-ALPES",
    "UNIVERS. GRENOBLE ALPES (EPE)": "UNIVERSITE GRENOBLE-ALPES",
    "UNIVERSITE GRENOBLE ALPES": "UNIVERSITE GRENOBLE-ALPES",
    "UNIV. POLYTECH. HAUTS France": "UNIVERSITE POLYTECHNIQUE DES HAUTS DE FRANCE",
    "UNIV. POLYTECH. HAUTS France (EPE)": "UNIVERSITE POLYTECHNIQUE DES HAUTS DE FRANCE",
    "UNIVERSITE  LORRAINE": "UNIVERSITE DE LORRAINE",
    "UNIVERSITE LORRAINE": "UNIVERSITE DE LORRAINE",
    "UNIVERSITE CHAMBERY": "UNIVERSITE SAVOIE MONT BLANC",
    "UNIVERSITE COTE D'AZUR": "UNIVERSITE COTE AZUR",
    "UNIVERSITE COTE D'AZUR (EPE)": "UNIVERSITE COTE AZUR",
    "UNIVERSITE DE LA GUYANE": "UNIVERSITE GUYANE",
    "UNIVERSITE DE LILLE (EPE)": "UNIVERSITE LILLE",
    "UNIVERSITE DES ANTILLES": "UNIVERSITE ANTILLES",
    "UNIVERSITE LYON 1": "UNIVERSITE LYON 1 (CLAUDE BERNARD)",
    "UNIVERSITE LYON 2": "UNIVERSITE LYON 2 (LUMIERE)",
    "UNIVERSITE LYON 3": "UNIVERSITE LYON 3 (JEAN MOULIN)",
    "UNIVERSITE MONTPELLIER (EPE)": "UNIVERSITE MONTPELLIER",
    "UNIVERSITE NANTES (EPE)": "UNIVERSITE NANTES",
    "UNIVERSITE PARIS 1": "UNIVERSITE PARIS 1 (PANTHEON SORBONNE)",
    "UNIVERSITE PARIS 2": "UNIVERSITE PARIS 2 (PANTHEON ASSAS)",
    "UNIVERSITE PARIS PANTHEON-ASSAS (PARIS 2)": "UNIVERSITE PARIS 2 (PANTHEON ASSAS)",
    "UNIVERSITE PARIS 3": "UNIVERSITE PARIS 3 (SORBONNE NOUVELLE)",
    "UNIVERSITE PARIS 8": "UNIVERSITE PARIS 8 (VINCENNES SAINT-DENIS)",
    "UNIVERSITE PARIS 13": "UNIVERSITE PARIS 13 (SORBONNE PARIS NORD)",
    "UNIVERSITE PARIS CITE (EPE)": "UNIVERSITE PARIS CITE",
    "UNIVERSITE PARIS-SACLAY": "UNIVERSITE PARIS SACLAY",
    "UNIVERSITE PARIS-SACLAY (EPE)": "UNIVERSITE PARIS SACLAY",
    "UNIVERSITE PARIS-EST \nCRETEIL": "UNIVERSITE PARIS-EST CRETEIL",
    "UNIVERSITE RENNES 1": "UNIVERSITE DE RENNES (RENNES 1)",
    "UNIVERSITE SORBONNE UNIVERSITE": "SORBONNE UNIVERSITE",
    "UNIVERSITE TOULOUSE 1": "UNIVERSITE TOULOUSE 1 (CAPITOLE)",
    "UNIVERSITE TOULOUSE CAPITOLE (TOULOUSE 1)": "UNIVERSITE TOULOUSE 1 (CAPITOLE)",
    "UNIVERSITE TOULOUSE 2": "UNIVERSITE TOULOUSE 2 (JEAN JAURES)",
    "UNIVERSITE TOULOUSE 3": "UNIVERSITE TOULOUSE 3 (PAUL SABATIER)",
    "UNIVERSITE VERSAILLES ST QUENT": "UNIVERSITE VERSAILLES ST QUENTIN",
    "UNIVERSITE CAEN": "UNIVERSITE DE CAEN NORMANDIE",
    "UNIVERSITE CLERMONT AUVERGNE": "UNIV. CLERMONT AUVERGNE (EPE)",
    "UNIVERSITE CLERMONT AUVERGNE (EPE)": "UNIV. CLERMONT AUVERGNE (EPE)",
    "UNIVERSITE GUSTAVE EIFFEL (EPE)": "UNIVERSITE GUSTAVE EIFFEL",
    "UNIVERSITE  NOUVELLE-CALEDONIE": "UNIVERSITE NOUVELLE-CALEDONIE",
    **SU.SIES_ALIASES,
}
# Groupings with no licence or BUT of their own in Parcoursup: kept as published, never linked.
UNLINKED = ("COMUE U BOURGOGNE-FRANCHE-COMT", "COMUE UNIV PARIS LUMIERES")


def canonical(name: str, aliases: dict[str, str] | None = None) -> str:
    """The B1 university name of a spelling (outer spaces only: "UNIVERSITE  BREST" keeps its two)."""
    table = NAME_ALIASES if aliases is None else aliases
    n = name if name in table else name.strip()
    return table.get(n, n)


# BUT formations without a Paysage id: one explicit rule per site (lien = "nom", graded B). The
# IUT de Toulouse (sites of Toulouse, Auch and Castres) is the IUT of Toulouse 3 - Paul Sabatier;
# its own id (1PpB9) is not in the Paysage register the warehouse loads, so it cannot be checked.
BUT_NAME_RULES = {
    r"^i\.?u\.?t\.? de toulouse\b": "UNIVERSITE TOULOUSE 3 (PAUL SABATIER)",
}
# Sites whose Paysage id names another establishment than the university the IUT belongs to: the
# rule wins over the id (lien = "nom", graded B). Allowed only where the id points to a university
# with no BUT row (red team P5 checks it), so a usable id is never overridden.
# The IUT de Nîmes is a component of the Université de Montpellier (confirmed by the owner); its
# Parcoursup formations carry the id of the Université de Nîmes, which has no IUT in the SIES table.
BUT_SITE_OVERRIDES = {
    r"^i\.?u\.?t\.? de nimes\b": "UNIVERSITE MONTPELLIER",
}

# Parcoursup BUT programme (fil_lib_voe_acc, folded, alphanumeric only, "parcours ..." removed)
# -> SIES specialty (Tableau 2). Explicit aliases for the spellings that differ.
SPECIALTY_ALIASES = {
    "science des donnees": "statistique et informatique decisionnelle",   # renamed in 2024
    "science et genie des materiaux": "sciences et genie des materiaux",
    "metiers de la transition et de l efficacite energetiques": "metiers de la transition et de l efficacite energetique",
}


def _key(label: str | None) -> str:
    t = re.sub(r"[^a-z0-9]+", " ", fold(label)).strip()
    return re.sub(r"\s+parcours\b.*$", "", t).strip()


def specialty_of(programme: str | None, known: set[str]) -> str | None:
    k = _key(programme)
    k = SPECIALTY_ALIASES.get(k, k)
    return k if k in known else None


# SIES général specialty codes (BUT Tableau 1) -> journey.SPECIALTIES codes
SPE_CODE = {"MATHS": "MATHS", "PHY_CHI": "PC", "HIST_GE": "HGGSP", "SES": "SES", "NUMERIQ": "NSI", "SVT": "SVT",
            "LLCER": "LLCER", "SC_INGE": "SI", "HLP": "HLP"}
MENTION_CODE = {"tres bien": "tb", "bien": "b", "assez bien": "ab", "passable 1er groupe": "passable1",
                "passable premier groupe": "passable1", "passable 2nd groupe": "passable2",
                "passable deuxieme groupe": "passable2", "inconnue": "inconnue"}
BAC_CODE = {"general": "bg", "technologique": "bt", "professionnel": "bp"}
# national master tables: their discipline label -> B1 discipline code
MASTER_DISC = {"droit, sciences politiques": "droit", "sc. eco, gestion, administration eco. et sociale": "eco",
               "arts, lettres, langues, sciences humaines et sociales": "alsh", "dont psychologie": "alsh:psychologie",
               "sciences-sante": "sciences", "sciences et techniques des activites physiques et sportives (staps)": "staps",
               "master enseignement": "meef", "ensemble": "toutes"}
# chain of the 2017 L1 entrants: L1 discipline (8 groups) -> B1 discipline when the match is one-to-one
CHAIN_DISC = {"droit, sciences politiques": "droit", "sc. eco, gestion, aes": "eco", "staps": "staps"}


def _num(v):
    return SU._num(v)


def _title(ws) -> str:
    for r in ws.iter_rows(min_row=1, max_row=1, values_only=True):
        return fold(str(r[0])) if r and r[0] else ""
    return ""


def detect(path: Path) -> str | None:
    """Which sprint-6 workbook this is, from its sheets (None for any other workbook)."""
    import openpyxl
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - not a readable workbook
        return None
    names = set(wb.sheetnames)
    t = {n: _title(wb[n]) for n in ("Tableau 1", "Tableau 2", "Tableau 3") if n in names}
    if "Réussite BUT établissements" in names:
        return "but_etab"
    if t.get("Tableau 2", "").startswith("tableau 2 - reussite au but"):
        return "but_nat"
    if t.get("Tableau 1", "").startswith("tableau 1 - reussite en 2 ou 3 ans par discipline et etablissement"):
        return "master_etab"
    if t.get("Tableau 3", "").startswith("tableau 3 - poursuite en master universitaire"):
        return "master_nat"
    if t.get("Tableau 2", "").startswith("parcours et reussite en licence et en master"):
        return "master_chain"
    return None


def find_sources(data_dir: Path | None = None) -> dict[str, Path]:
    out = {}
    for p in sorted((data_dir or config.DATA_DIR / "sources").glob("*.xlsx")):
        kind = detect(p)
        if kind:
            out[kind] = p             # sorted names: the latest note number wins
    return out


def _rows(ws):
    return [r for r in ws.iter_rows(values_only=True)]


def _cols(head, spec, where):
    h = [fold(str(x).replace("\n", " ")) if x is not None else "" for x in head]
    out = {}
    for key, pred in spec.items():
        idx = [i for i, x in enumerate(h) if pred(x)]
        if len(idx) != 1:
            raise ValueError(f"{where}: column {key} found {len(idx)} times in {h}")
        out[key] = idx[0]
    return out


def _cohort(title: str, rx: str) -> int | None:
    m = re.search(rx, fold(title))
    return int(m.group(1)) if m else None


def _record(r, cols):
    rec, masked = {}, {}
    for key, i in cols.items():
        v, why = _num(r[i] if i < len(r) else None)
        rec[key] = v
        if why:
            masked[key] = why
    rec["masked"] = masked
    return rec


# ------------------------------------------------------------------------- BUT
BUT_COLUMNS = {
    "reussite3": {"inscrits": lambda h: h.startswith("premieres inscriptions"),
                  "n_reussite": lambda h: h == "reussite en 3 ans",
                  "taux": lambda h: h == "reussite en 3 ans taux observe",
                  "taux_simule": lambda h: h == "reussite en 3 ans taux simule",
                  "va": lambda h: h.startswith("reussite en 3 ans valeur ajoutee")},
    "poursuite": {"n_diplomes": lambda h: h == "effectif",
                  "taux_poursuite": lambda h: h.startswith("poursuite d'etudes"),
                  "taux_master": lambda h: h.startswith("dont master"),
                  "taux_commerce": lambda h: h.startswith("dont ecole de commerce"),
                  "taux_ingenieur": lambda h: h.startswith("dont ecole d'ingenieur")},
}
BUT_SHEETS = {"reussite3": "Réussite BUT établissements", "poursuite": "Poursuite diplômés 3ans"}


def parse_but_etab(path: Path, aliases: dict[str, str] | None = None) -> list[dict]:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    for kind, sheet in BUT_SHEETS.items():
        rows = _rows(wb[sheet])
        hi = next(i for i, r in enumerate(rows) if r and r[0] and fold(str(r[0])).startswith("nom etablissement"))
        cols = _cols(rows[hi], BUT_COLUMNS[kind], sheet)
        cohort = _cohort(str(rows[0][0]), r"bacheliers (\d{4})")
        for r in rows[hi + 1:]:
            if not r or not r[0] or r[1] is None:
                continue
            name = str(r[0])
            if name == "Etudiants":
                continue              # a person counted once: does not add up, not the reference
            etab = NATIONAL if name in ("Inscriptions principales", "Ensemble") else canonical(name, aliases)
            out.append({"kind": kind, "cohort": cohort, "etab": etab, "etab_row": name, **_record(r, cols)})
    return out


def parse_but_nat(path: Path) -> list[dict]:
    """National BUT rates: by bac series / général doublette / mention (Tableau 1), by specialty (Tableau 2)."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    spec = {"part": lambda h: h == "part des inscrits", "passage2": lambda h: h == "taux de passage en 2eme annee",
            "passage3": lambda h: h == "taux de passage en 3eme annee", "taux": lambda h: h == "taux de reussite en 3 ans"}
    out = []
    for sheet, table in (("Tableau 1", "profil"), ("Tableau 2", "specialite")):
        rows = _rows(wb[sheet])
        hi = next(i for i, r in enumerate(rows) if r and r[1] is not None and fold(str(r[1])) == "part des inscrits")
        cols = _cols(rows[hi], spec, sheet)
        cohort = _cohort(str(rows[0][0]), r"bacheliers (\d{4})")
        section = "serie" if table == "profil" else "specialite"
        for r in rows[hi + 1:]:
            if not r or not r[0]:
                continue
            label = re.sub(r"\s+", " ", str(r[0])).strip()
            f = fold(label)
            if r[1] is None:
                if f.startswith("mention"):
                    section = "mention"
                continue
            if f == "ensemble":
                code, dim = "ensemble", "ensemble"
            elif table == "specialite":
                dim = "secteur" if f.startswith("secteur") else "specialite"
                code = _key(label)
            elif section == "mention":
                dim, code = "mention", MENTION_CODE.get(f)
            else:
                dim = "serie"
                m = re.match(r"dont ([A-Z_]+) / ([A-Z_]+)$", label)
                if m:
                    a, b = SPE_CODE.get(m.group(1)), SPE_CODE.get(m.group(2))
                    code = f"bg:{'+'.join(sorted((a, b)))}" if a and b else None
                elif f.startswith("ensemble bac general"):
                    code = "bg"
                elif f.startswith("ensemble bac technologique"):
                    code = "bt"
                elif f.startswith("dont stmg"):
                    code = "bt:stmg"
                elif f.startswith("dont autres bac techno"):
                    code = "bt:autres"
                elif f.startswith("ensemble bac professionnel"):
                    code = "bp"
                else:
                    code = None
            if code is None:
                raise ValueError(f"{sheet}: unrecognised row {label!r}")
            out.append({"tab": table, "dim": dim, "code": code, "label": label, "cohort": cohort, **_record(r, cols)})
    return out


# ------------------------------------------------------------------------- master
MASTER_COLUMNS = {
    "reussite23": {"inscrits": lambda h: h.startswith("effectif"), "n_reussite": lambda h: h == "reussite a 2 ou 3 ans",
                   "taux": lambda h: h == "taux observe", "taux_simule": lambda h: h == "taux simule",
                   "va": lambda h: h.startswith("valeur ajoutee")},
    "passage_m2": {"inscrits": lambda h: h == "premieres inscriptions en m1", "n_passage": lambda h: h == "passage en m2",
                   "taux": lambda h: h == "passage en m2 taux observe", "taux_simule": lambda h: h == "passage en m2 taux simule",
                   "va": lambda h: h.startswith("passage en m2 valeur ajoutee"),
                   "taux_redoublement": lambda h: h == "taux de redoublement en m1",
                   "taux_reo_univ": lambda h: h == "taux de reorientation a l'universite",
                   "taux_reo_ingenieur": lambda h: h == "taux de reorientation en ecole d'ingenieur ou ens",
                   "taux_reo_management": lambda h: h == "taux de reorientation en ecole de management",
                   "taux_reo_autres": lambda h: h == "taux autres reorientations",
                   "taux_sortie": lambda h: h == "taux de sortie"},
}
MASTER_SHEETS = {"reussite23": ("Tableau 1", r"m1 en (\d{4})"), "passage_m2": ("Tableau 3", r"m1 en (\d{4})")}
DEVENIR_M2 = ("taux", "taux_redoublement", "taux_reo_univ", "taux_reo_ingenieur", "taux_reo_management", "taux_reo_autres", "taux_sortie")


def parse_master_etab(path: Path, aliases: dict[str, str] | None = None) -> list[dict]:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    for kind, (sheet, rx) in MASTER_SHEETS.items():
        rows = _rows(wb[sheet])
        hi = next(i for i, r in enumerate(rows) if r and r[0] and fold(str(r[0])).startswith("discipline"))
        cols = _cols(rows[hi], MASTER_COLUMNS[kind], sheet)
        cohort = _cohort(str(rows[0][0]), rx)
        block = None
        for r in rows[hi + 1:]:
            if not r or not r[0] or not r[1] or r[2] is None:
                continue
            disc = SU.DISC_CODE.get(fold(str(r[0]).replace("\n", " ")))
            if disc is None:
                continue                  # notes and footers
            name = str(r[1])
            if disc == "toutes":
                block = name
            if block == "Etudiants":
                continue
            etab = NATIONAL if block == "Inscriptions principales" else canonical(block, aliases)
            out.append({"kind": kind, "cohort": cohort, "etab": etab, "etab_row": name, "block_row": block,
                        "discipline": disc, **_record(r, cols)})
    return out


def parse_master_nat(path: Path) -> list[dict]:
    """Continuation into master of licence graduates, by discipline (Tableau 3), 2021-2024."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = _rows(wb["Tableau 3"])
    hi = next(i for i, r in enumerate(rows) if r and r[0] and fold(str(r[0])) == "discipline de lg")
    years = {i: int(m.group(1)) for i, h in enumerate(rows[hi + 1]) if h and (m := re.search(r"(\d{4})", str(h)))}
    out = []
    for r in rows[hi + 2:]:
        if not r or not r[0] or r[1] is None:
            continue
        code = MASTER_DISC.get(fold(str(r[0])))
        if code is None:
            raise ValueError(f"Tableau 3: unrecognised discipline {r[0]!r}")
        for i, y in years.items():
            v, why = _num(r[i] if i < len(r) else None)
            out.append({"tab": "poursuite_lg", "dim": "discipline", "code": code, "label": str(r[0]), "cohort": y,
                        "taux": v, "masked": {"taux": why} if why else {}})
    return out


def parse_master_chain(path: Path) -> list[dict]:
    """The 2017 L1 entrants followed to master (Tableau 2), by bac type, mention and L1 discipline."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = _rows(wb["Tableau 2"])
    hi = next(i for i, r in enumerate(rows) if r and r[0] and fold(str(r[0])).startswith("caracteristiques"))
    spec = {"part": lambda h: h == "part des inscrits", "licence34": lambda h: h == "reussite en licence en 3 ou 4 ans",
            "passage_master": lambda h: h.startswith("passage en master"),
            "master56": lambda h: h == "reussite au master apres 5 ou 6 ans d'etudes",
            "master567": lambda h: h == "reussite au master apres 5, 6 ans ou 7 ans d'etudes"}
    cols = _cols(rows[hi], spec, "Tableau 2")
    cohort = _cohort(str(rows[0][0]), r"bacheliers (\d{4})")
    out, section = [], "sexe"
    for r in rows[hi + 1:]:
        if not r or not r[0]:
            continue
        f = fold(str(r[0]))
        if r[1] is None:
            section = ("bac" if f.startswith("type de bac") else "mention" if f.startswith("mention")
                       else "discipline" if f.startswith("discipline") else section)
            continue
        if f == "ensemble":
            dim, code = "ensemble", "ensemble"
        elif section == "bac":
            dim, code = "bac", BAC_CODE.get(f)
        elif section == "mention":
            dim, code = "mention", MENTION_CODE.get(f)
        elif section == "discipline":
            dim, code = "discipline", CHAIN_DISC.get(f, "l1:" + _key(str(r[0])))
        else:
            dim, code = "sexe", f
        if code is None:
            raise ValueError(f"Tableau 2: unrecognised row {r[0]!r}")
        out.append({"tab": "chaine_l1", "dim": dim, "code": code, "label": str(r[0]), "cohort": cohort, **_record(r, cols)})
    return out


# ------------------------------------------------------------------------- warehouse
NAT_KEYS = ("tab", "dim", "code", "label", "cohort", "part", "passage2", "passage3", "taux", "licence34",
            "passage_master", "master56", "master567")
BUT_KEYS = ("kind", "cohort", "etab", "etab_row", "inscrits", "n_reussite", "taux", "taux_simule", "va",
            "n_diplomes", "taux_poursuite", "taux_master", "taux_commerce", "taux_ingenieur")
MASTER_KEYS = ("kind", "cohort", "etab", "etab_row", "block_row", "discipline", "inscrits", "n_reussite", "n_passage",
               "taux", "taux_simule", "va", "taux_redoublement", "taux_reo_univ", "taux_reo_ingenieur",
               "taux_reo_management", "taux_reo_autres", "taux_sortie")


def build(con, sources: dict[str, Path] | None = None, paysage_table: dict[str, str] | None = None,
          name_rules: dict[str, str] | None = None, aliases: dict[str, str] | None = None,
          overrides: dict[str, str] | None = None) -> dict:
    """core.but_success, core.master_success, core.poursuite_national, core.but_success_link.
    Returns counts (empty tables without sources)."""
    src = find_sources() if sources is None else sources
    con.execute("CREATE SCHEMA IF NOT EXISTS core")
    con.execute(f"""CREATE OR REPLACE TABLE core.but_success ({", ".join(f"{k} {'VARCHAR' if k in ('kind', 'etab', 'etab_row') else 'DOUBLE'}" for k in BUT_KEYS)},
        masked VARCHAR, source VARCHAR)""")
    con.execute(f"""CREATE OR REPLACE TABLE core.master_success ({", ".join(f"{k} {'VARCHAR' if k in ('kind', 'etab', 'etab_row', 'block_row', 'discipline') else 'DOUBLE'}" for k in MASTER_KEYS)},
        masked VARCHAR, source VARCHAR)""")
    con.execute(f"""CREATE OR REPLACE TABLE core.poursuite_national (source_kind VARCHAR, {", ".join(f"{k} {'VARCHAR' if k in ('tab', 'dim', 'code', 'label') else 'DOUBLE'}" for k in NAT_KEYS)},
        masked VARCHAR, source VARCHAR)""")
    con.execute("""CREATE OR REPLACE TABLE core.but_success_link (session INTEGER, cod_aff_form VARCHAR, etab VARCHAR,
        specialite VARCHAR, programme VARCHAR, lien VARCHAR, id_paysage VARCHAR)""")
    counts = {}

    def ins(table, keys, recs, extra):
        con.executemany(f"INSERT INTO {table} VALUES ({', '.join('?' * (len(keys) + len(extra(recs[0])) + 2))})",
                        [extra(r) + [r.get(k) for k in keys] + [json.dumps(r["masked"], sort_keys=True), r["_src"]] for r in recs]) if recs else None

    def load(kind, parser):
        if kind not in src:
            return []
        recs = parser(Path(src[kind]), aliases) if kind in ("but_etab", "master_etab") else parser(Path(src[kind]))
        for r in recs:
            r["_src"] = Path(src[kind]).name
        counts[kind] = len(recs)
        return recs

    but = load("but_etab", parse_but_etab)
    ins("core.but_success", BUT_KEYS, but, lambda r: [])
    master = load("master_etab", parse_master_etab)
    ins("core.master_success", MASTER_KEYS, master, lambda r: [])
    for kind, parser in (("but_nat", parse_but_nat), ("master_nat", parse_master_nat), ("master_chain", parse_master_chain)):
        recs = load(kind, parser)
        ins("core.poursuite_national", NAT_KEYS, recs, lambda r, k=kind: [k])
    if not but:
        return counts
    # BUT formations -> university (Paysage id first, then the explicit site rules) and specialty
    latest = con.execute("SELECT max(session) FROM core.formation_year").fetchone()[0]
    has_ref = con.execute("""SELECT count(*) FROM information_schema.tables WHERE table_schema = 'ref'
                             AND table_name = 'formation_etablissement'""").fetchone()[0]
    ids = dict(con.execute("""SELECT cod_aff_form, id_resolu FROM ref.formation_etablissement
                              WHERE session = ? AND id_resolu IS NOT NULL""", [latest]).fetchall()) if has_ref else {}
    known_etab = {r["etab"] for r in but}
    known_spe = {r[0] for r in con.execute("""SELECT code FROM core.poursuite_national
                                              WHERE source_kind = 'but_nat' AND dim = 'specialite'""").fetchall()}
    by_id = paysage_table or SU.PAYSAGE_SIES
    rules = name_rules or BUT_NAME_RULES
    site = BUT_SITE_OVERRIDES if overrides is None else overrides
    links = []
    for cod, etab, prog, contrat in con.execute("""SELECT cod_aff_form, g_ea_lib_vx, fil_lib_voe_acc, contrat_etab
            FROM core.formation_year WHERE session = ? AND fili = 'BUT'""", [latest]).fetchall():
        pid = ids.get(cod)
        uni, lien = None, None
        forced = [u for rx, u in site.items() if re.search(rx, fold(etab))]
        if contrat == "Public":
            if len(forced) == 1 and by_id.get(pid) not in known_etab:
                uni, lien = forced[0], "nom"
            elif pid is not None:
                uni, lien = by_id.get(pid), "identifiant"
            else:
                hits = [u for rx, u in rules.items() if re.search(rx, fold(etab))]
                if len(hits) == 1:
                    uni, lien = hits[0], "nom"
            if uni not in known_etab:
                uni, lien = None, None
        links.append([latest, cod, uni, specialty_of(prog, known_spe), prog, lien, pid])
    con.executemany("INSERT INTO core.but_success_link VALUES (?, ?, ?, ?, ?, ?, ?)", links)
    counts["but_links"] = sum(1 for x in links if x[2])
    return counts


# ------------------------------------------------------------------------- serving
def _fetch(con, sql, args):
    cur = con.execute(sql, args)
    cols = [d[0] for d in cur.description]
    out = []
    for r in cur.fetchall():
        d = dict(zip(cols, r))
        d["masked"] = json.loads(d.get("masked") or "{}")
        if d.get("cohort") is not None:
            d["cohort"] = int(d["cohort"])
        if "inscrits" in d:
            d["small"] = d["inscrits"] is not None and d["inscrits"] < SMALL_COHORT
        out.append(d)
    return out


def _national(con, kind, table):
    return _fetch(con, """SELECT * FROM core.poursuite_national WHERE source_kind = ? AND tab = ?
                          ORDER BY dim, code, cohort""", [kind, table])


def _tables_ready(con, *tables) -> bool:
    try:
        return all(con.execute(f"SELECT count(*) FROM {t}").fetchone()[0] > 0 for t in tables)
    except Exception:  # noqa: BLE001 - warehouse built before sprint 6
        return False


def but_for_formation(con, cod: str, fili: str | None, contrat: str | None) -> dict | None:
    """The BUT block of a formation page, or None when the formation is not a BUT."""
    if fili != "BUT" or not _tables_ready(con, "core.but_success"):
        return None
    link = con.execute("""SELECT etab, specialite, programme, lien FROM core.but_success_link
                          WHERE cod_aff_form = ? ORDER BY session DESC LIMIT 1""", [cod]).fetchone()
    nat_rows = _national(con, "but_nat", "profil") + _national(con, "but_nat", "specialite")
    nat = {k: _fetch(con, "SELECT * FROM core.but_success WHERE kind = ? AND etab = ?", [k, NATIONAL]) for k in BUT_SHEETS}
    spe = link[1] if link else None
    spe_row = next((r for r in nat_rows if r["tab"] == "specialite" and r["code"] == spe), None)
    out = {"status": "ok", "etab": None, "etab_label": None, "kinds": {},
           "national": {k: (v[0] if v else None) for k, v in nat.items()},
           "specialite": spe_row, "specialite_programme": link[2] if link else None,
           "profil": [r for r in nat_rows if r["tab"] == "profil"], "lien": None}
    if contrat != "Public":
        out.update(status="hors_champ", reason="le SIES publie ces taux pour les IUT des universités publiques seulement")
    elif link is None or link[0] is None:
        out.update(status="non_rattachee", reason="cet IUT n'a pas pu être rattaché avec certitude à une université des tableaux SIES")
    else:
        etab, _, _, lien = link
        out.update(etab=etab, etab_label=SU.UNIVERSITY_LABEL.get(etab, etab), lien=lien)
        for k in BUT_SHEETS:
            u = _fetch(con, "SELECT * FROM core.but_success WHERE kind = ? AND etab = ?", [k, etab])
            if u:
                out["kinds"][k] = u[0]
    src = SOURCE_LABEL["but"]
    r = out["kinds"].get("reussite3")
    if r:
        n = out["national"].get("reussite3") or {}
        r["maillon"] = E.maillon("but_3_ans", r["taux"], france=n.get("taux"), attendu=r["taux_simule"],
                                 granularite="universite", population="néo-bacheliers inscrits en 1re année de BUT",
                                 cohorte=r["cohort"], effectif=r["inscrits"], source=src, lien=out["lien"],
                                 avertissement="tous les BUT de l'université, toutes spécialités, pas seulement celui-ci")
    p = out["kinds"].get("poursuite")
    if p:
        n = out["national"].get("poursuite") or {}
        p["maillon"] = E.maillon("poursuite_apres_but", p["taux_poursuite"], france=n.get("taux_poursuite"),
                                 granularite="universite", population="diplômés du BUT en 3 ans",
                                 cohorte=p["cohort"], effectif=p["n_diplomes"], source=src, lien=out["lien"],
                                 avertissement="poursuite d'études un an après le diplôme, toutes spécialités de l'université")
    if out["status"] == "ok" and not out["kinds"]:
        out.update(status="non_publie", reason="le SIES ne publie pas de taux pour cette université")
    return out


def master_for_formation(con, cod: str, fili: str | None, contrat: str | None) -> dict | None:
    """The "after the licence: master" block of a licence page (same university and discipline as B1)."""
    if fili not in SU.LICENCE_FILI or not _tables_ready(con, "core.master_success"):
        return None
    try:
        link = con.execute("""SELECT etab, discipline, lien FROM core.licence_success_link
                              WHERE cod_aff_form = ? ORDER BY session DESC LIMIT 1""", [cod]).fetchone()
    except Exception:  # noqa: BLE001 - no B1 link table
        link = None
    disc = link[1] if link else None
    chain = _national(con, "master_chain", "chaine_l1")
    pours = [r for r in _national(con, "master_nat", "poursuite_lg") if r["code"] in ((disc or "toutes"), "toutes")]
    out = {"status": "ok", "etab": None, "etab_label": None, "discipline": disc,
           "discipline_label": SU.DISCIPLINES[disc][1] if disc else None, "kinds": {}, "national": {},
           "poursuite_lg": pours, "chaine": chain, "lien": None}
    for k in MASTER_SHEETS:
        n = _fetch(con, "SELECT * FROM core.master_success WHERE kind = ? AND etab = ? AND discipline = ?",
                   [k, NATIONAL, disc or "toutes"])
        if n:
            out["national"][k] = n[0]
    if contrat != "Public":
        out.update(status="hors_champ", reason="le SIES publie ces taux pour les universités publiques seulement")
    elif link is None or link[0] is None:
        out.update(status="non_rattachee", reason="cet établissement n'a pas pu être rattaché avec certitude à une université des tableaux SIES")
    elif disc is None:
        out.update(status="discipline_non_classee", etab=link[0], etab_label=SU.UNIVERSITY_LABEL.get(link[0], link[0]),
                   reason="cette licence ne relève pas clairement d'une seule des cinq grandes disciplines du SIES")
    else:
        etab, lien = link[0], link[2] or "nom"
        out.update(etab=etab, etab_label=SU.UNIVERSITY_LABEL.get(etab, etab), lien=lien)
        for k in MASTER_SHEETS:
            u = _fetch(con, "SELECT * FROM core.master_success WHERE kind = ? AND etab = ? AND discipline = ?", [k, etab, disc])
            if u:
                out["kinds"][k] = u[0]
        if not out["kinds"]:
            out.update(status="non_publie", reason="le SIES ne publie pas de master dans cette discipline pour cette université")
    etape = {"reussite23": "master_2_ou_3_ans", "passage_m2": "passage_m2"}
    for k, u in out["kinds"].items():
        n = out["national"].get(k) or {}
        u["maillon"] = E.maillon(etape[k], u["taux"], france=n.get("taux"), attendu=u["taux_simule"],
                                 granularite="universite_x_discipline", population="étudiants inscrits pour la première fois en M1",
                                 cohorte=u["cohort"], effectif=u["inscrits"], source=SOURCE_LABEL["master"], lien=out["lien"],
                                 avertissement="tous les entrants en M1 de cette discipline dans l'université, d'où qu'ils viennent")
    return out
