"""Sprint 7: gateways and reorientations.

Two sources, checked on the real data before the design:

1. The "Réorientations - cohorte 2023" sheet of the B1 workbook (SIES NF 2025-29): where the
   2023 bacheliers entering L1 who changed course in 2024 went, per university x broad discipline.
   * Every rate is a share of the L1 ENTRANTS (not of the reoriented students).
   * Reoriented = hors université + à l'université (exact for every university). The published
     detail is NOT exhaustive: hors université lists CPGE, STS (apprenticeship or not), business
     and engineering schools; à l'université lists BUT and the nursing state diploma. The rest is
     not detailed and is never computed by difference here.
   * Masked ("ns") below 100 entrants; universities add up to "Inscriptions principales" (on that
     national row, hors + université exceeds the total by 1-3 students: multiple enrolments).
   * The rows are grouped by block, as in B1, and linked through B1's own university link.

2. MonMaster (SIES open data, 2024 and 2025 sessions; 2023 has another layout and no
   "admitted from the same establishment" field). One row per master offer (`ifc`: programme x
   site x apprenticeship), with applications, proposals and acceptances by origin. Checked:
   origins add up to the totals, main + complementary phase = total, the funnel never grows,
   published percentages = counts, nothing masked. Units: applications summed over programmes
   are VŒUX (one person applies to several masters); acceptances are people. No proposal /
   application ratio is shown: it would read as a chance of admission.
   Masters are linked to a licence through the university of B1's link (its Paysage id) and the
   broad discipline (MASTER_DISCIPLINE). MEEF (teaching) masters recruit from every licence and
   stay out of the discipline figures; mixed or theology masters get no discipline.

The ministry figures are stored and shown as published, never recomputed.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

from . import config
from . import evidence as E
from . import success as SU
from .success import fold

NATIONAL = SU.NATIONAL
MONMASTER = ("monmaster_2024_2025", "monmaster_2023")
MASTER_SESSIONS = ("2024", "2025")
SMALL = 100

# ------------------------------------------------------------------------- reorientation sheet
REO_TITLE = "detail des formations suivies"
# key -> (header predicate on the folded header, family) ; family None = the totals
REO_COLUMNS = {
    "inscrits": (lambda h: h.startswith("inscrits en l1"), None),
    "n_reorientation": (lambda h: h == "reorientation (rappel)", None),
    "taux_reorientation": (lambda h: h == "taux de reorientation (rappel)", None),
    "taux_hors_univ": (lambda h: h == "taux de reorientation hors universite", None),
    "n_hors_univ": (lambda h: h == "dont reorientation hors universite", None),
    "taux_univ": (lambda h: h == "taux de reorientation a l'universite", None),
    "n_univ": (lambda h: h == "dont reorientation a l'universite", None),
    "taux_cpge": (lambda h: h == "taux de reorientation en cpge", "hors"),
    "taux_sts_app": (lambda h: h == "taux de reorientation en sts en apprentissage", "hors"),
    "taux_sts": (lambda h: h == "taux de reorientation en sts hors apprentissage", "hors"),
    "taux_commerce": (lambda h: h == "taux de reorientation en ecole de commerce", "hors"),
    "taux_ingenieur": (lambda h: h == "taux de reorientation en ecole d'ingenieur", "hors"),
    "taux_but": (lambda h: h == "taux de reorientation en dut/but", "univ"),
    "taux_infirmier": (lambda h: h == "taux de reorientation en de infirmier", "univ"),
}
REO_DETAIL = [k for k, (_, fam) in REO_COLUMNS.items() if fam]
# spellings of this sheet that differ from B1's university names (B1's SIES_ALIASES also apply)
REO_ALIASES = {"UNIVERSITE PARIS 13 SORBONNE PARIS NORD": "UNIVERSITE PARIS 13 (SORBONNE PARIS NORD)"}


def parse_reorientation(path: Path, aliases: dict[str, str] | None = None) -> list[dict]:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = next((w for w in wb.worksheets if fold(str(next(w.iter_rows(max_row=1, values_only=True))[0] or "")).startswith(REO_TITLE)), None)
    if ws is None:
        return []
    rows = list(ws.iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and r[0] and fold(str(r[0])).startswith("nom etablissement"))
    head = [fold(str(h).replace("\n", " ")) if h else "" for h in rows[hi]]
    cols = {}
    for key, (pred, _) in REO_COLUMNS.items():
        idx = [i for i, h in enumerate(head) if pred(h)]
        if len(idx) != 1:
            raise ValueError(f"reorientation sheet: column {key} found {len(idx)} times in {head}")
        cols[key] = idx[0]
    m = re.search(r"rentree (\d{4})", fold(str(rows[0][0])))
    cohort = int(m.group(1)) if m else None
    out, block = [], None
    for r in rows[hi + 1:]:
        if not r or not r[0] or not r[1]:
            continue
        code = SU.DISC_CODE.get(fold(str(r[1]).replace("\n", " ")))
        if code is None:
            continue
        name = str(r[0]).strip()
        if code == "toutes":
            block = (REO_ALIASES if aliases is None else aliases).get(name, SU.SIES_ALIASES.get(name, name))
        if block in SU.SKIPPED_ROWS:
            continue
        rec = {"cohort": cohort, "etab": NATIONAL if block == SU.NATIONAL_ROW else block, "etab_row": name,
               "discipline": code, "masked": {}}
        for key, i in cols.items():
            v, why = SU._num(r[i] if i < len(r) else None)
            rec[key] = v
            if why:
                rec["masked"][key] = why
        out.append(rec)
    return out


# ------------------------------------------------------------------------- MonMaster
# MonMaster broad discipline (disci_lib) -> B1 discipline of the licence. None: no single one.
MASTER_DISCIPLINE = {
    "droit, sciences politiques": "droit",
    "sciences economiques, gestion": "eco",
    "administration economique et sociale": "eco",
    "langues": "alsh",
    "lettres, sciences du langage, arts": "alsh",
    "pluridisciplinaire lettres, langues, sciences humaines": "alsh",
    "sciences humaines et sociales": "alsh",
    "sciences fondamentales et applications": "sciences",
    "sciences de la vie, de la terre et de l'univers": "sciences",
    "pluridisciplinaire sciences": "sciences",
    "staps": "staps",
    "pluridisciplinaire droit, sciences economiques, aes": None,   # droit or eco: not guessed
    "theologie": None,
}
ORIGINS = ("lg3", "lp3", "but3", "master", "autre", "noninscri")
ORIGIN_LABEL = {"lg3": "L3 de licence générale", "lp3": "licence professionnelle", "but3": "BUT 3e année",
                "master": "déjà en master", "autre": "autre formation", "noninscri": "non inscrits l'an passé"}


def is_meef(disci_master: str | None) -> bool:
    return fold(disci_master).replace(".", "") == "meef"


def master_discipline(disci_lib: str | None, disci_master: str | None) -> str | None:
    if is_meef(disci_master):
        return "meef"
    return MASTER_DISCIPLINE.get(fold(disci_lib))


def build(con, raw_dir: Path = config.RAW_DIR, reo_source: Path | None | object = ..., aliases: dict[str, str] | None = None) -> dict:
    """ref.monmaster_raw, core.licence_reorientation, core.monmaster. Returns counts."""
    out = {"monmaster_raw": build_raw(con, raw_dir)}
    con.execute(f"""CREATE OR REPLACE TABLE core.licence_reorientation (cohort INTEGER, etab VARCHAR, etab_row VARCHAR,
        discipline VARCHAR, {", ".join(f"{k} DOUBLE" for k in REO_COLUMNS)}, masked VARCHAR, source VARCHAR)""")
    src = SU.source_path() if reo_source is ... else reo_source
    if src is not None and Path(src).exists():
        recs = parse_reorientation(Path(src), aliases)
        keys = list(REO_COLUMNS)
        con.executemany(f"INSERT INTO core.licence_reorientation VALUES ({', '.join('?' * (len(keys) + 6))})",
                        [[r["cohort"], r["etab"], r["etab_row"], r["discipline"], *[r[k] for k in keys],
                          json.dumps(r["masked"], sort_keys=True), Path(src).name] for r in recs])
        out["reorientation"] = len(recs)
    out["monmaster"] = build_monmaster(con)
    return out


def build_raw(con, raw_dir: Path = config.RAW_DIR) -> int:
    """ref.monmaster_raw: every MonMaster row as published (union by column name); empty when absent."""
    con.execute("CREATE SCHEMA IF NOT EXISTS ref")
    con.execute("CREATE SCHEMA IF NOT EXISTS core")
    paths = [(k, Path(raw_dir) / f"{k}.parquet") for k in MONMASTER]
    paths = [(k, p) for k, p in paths if p.exists()]
    if not paths:
        # no raw files (development copy of a committed warehouse): keep the staged table if any
        con.execute("CREATE TABLE IF NOT EXISTS ref.monmaster_raw (_source VARCHAR)")
        return con.execute("SELECT count(*) FROM ref.monmaster_raw").fetchone()[0]
    union = " UNION ALL BY NAME ".join(f"SELECT '{k}' AS _source, * FROM read_parquet('{p}')" for k, p in paths)
    con.execute(f"CREATE OR REPLACE TABLE ref.monmaster_raw AS {union}")
    return con.execute("SELECT count(*) FROM ref.monmaster_raw").fetchone()[0]


MM_COLS = ("session", "ifc", "id_paysage", "eta_nom", "inm", "mention", "parcours", "alternance", "disci_lib", "disci_master",
           "discipline", "n_can", "n_prop", "n_accept", "n_accept_etab", *[f"n_accept_{o}" for o in ORIGINS])


def build_monmaster(con) -> int:
    con.execute(f"""CREATE OR REPLACE TABLE core.monmaster ({", ".join(f"{c} {'BIGINT' if c.startswith('n_') else 'VARCHAR'}" for c in MM_COLS)})""")
    have = {r[0] for r in con.execute("DESCRIBE ref.monmaster_raw").fetchall()}
    if "n_can_pp" not in have:
        return 0
    # the B1 discipline of each distinct (disci_lib, disci_master) pair, then one set-based insert
    pairs = con.execute(f"""SELECT DISTINCT disci_lib, disci_master FROM ref.monmaster_raw WHERE session IN {MASTER_SESSIONS}""").fetchall()
    con.execute("CREATE OR REPLACE TEMP TABLE _mm_disc (disci_lib VARCHAR, disci_master VARCHAR, discipline VARCHAR)")
    con.executemany("INSERT INTO _mm_disc VALUES (?, ?, ?)", [(a, b, master_discipline(a, b)) for a, b in pairs])
    con.execute(f"""INSERT INTO core.monmaster
        SELECT r.session, r.ifc, coalesce(r.id_paysage_actuel, r.id_paysage), r.eta_nom, r.inm, r.mention, r.parcours, r.alternance,
               r.disci_lib, r.disci_master, d.discipline, r.n_can_pp + r.n_can_pc, r.n_prop_total, r.n_accept_total, r.n_accept_etab,
               {", ".join(f"r.n_accept_{o}_total" for o in ORIGINS)}
        FROM ref.monmaster_raw r JOIN _mm_disc d ON d.disci_lib IS NOT DISTINCT FROM r.disci_lib
                                               AND d.disci_master IS NOT DISTINCT FROM r.disci_master
        WHERE r.session IN {MASTER_SESSIONS}""")
    return con.execute("SELECT count(*) FROM core.monmaster").fetchone()[0]


# ------------------------------------------------------------------------- serving
def _ready(con, table) -> bool:
    try:
        return con.execute(f"SELECT count(*) FROM {table}").fetchone()[0] > 0
    except Exception:  # noqa: BLE001 - warehouse built before sprint 7
        return False


def _reo_row(con, etab, disc):
    cur = con.execute("SELECT * FROM core.licence_reorientation WHERE etab = ? AND discipline = ?", [etab, disc])
    r = cur.fetchone()
    if r is None:
        return None
    d = dict(zip([c[0] for c in cur.description], r))
    d["masked"] = json.loads(d["masked"] or "{}")
    d["small"] = d["inscrits"] is not None and d["inscrits"] < SMALL
    return d


def masters_of(con, pid: str | None, disc: str | None, session: str) -> dict | None:
    """Aggregate of the masters of one establishment (Paysage id) in one B1 discipline (None: France)."""
    where, args = ["session = ?", "discipline = ?"], [session, disc]
    if pid is not None:
        where.append("id_paysage = ?")
        args.append(pid)
    w = " AND ".join(where)
    tot = con.execute(f"""SELECT count(*), count(DISTINCT inm), sum(n_can), sum(n_prop), sum(n_accept), sum(n_accept_etab),
            {", ".join(f"sum(n_accept_{o})" for o in ORIGINS)} FROM core.monmaster WHERE {w}""", args).fetchone()
    if not tot[0]:
        return None
    n_acc = tot[4] or 0
    agg = {"session": int(session), "n_offres": tot[0], "n_mentions": tot[1], "n_can": tot[2], "n_prop": tot[3],
           "n_accept": n_acc, "n_accept_etab": tot[5],
           "origines": {o: tot[6 + i] for i, o in enumerate(ORIGINS)},
           # shares of a SUM of published counts: labelled "calculé" on the page
           "part_etab": (100 * tot[5] / n_acc) if n_acc else None,
           "parts_origine": {o: (100 * tot[6 + i] / n_acc) if n_acc else None for i, o in enumerate(ORIGINS)}}
    if pid is not None:
        agg["mentions"] = [dict(zip(("mention", "n_offres", "n_can", "n_accept", "n_accept_etab"), r)) for r in con.execute(
            f"""SELECT mention, count(*), sum(n_can), sum(n_accept), sum(n_accept_etab) FROM core.monmaster WHERE {w}
                GROUP BY mention ORDER BY sum(n_accept) DESC, mention LIMIT 8""", args).fetchall()]
    return agg


def latest_session(con) -> str | None:
    r = con.execute("SELECT max(session) FROM core.monmaster").fetchone()
    return r[0] if r else None


def for_licence(con, cod: str, fili: str | None, contrat: str | None, paysage: dict[str, str] | None = None) -> dict | None:
    """The gateways block of a licence page: reorientation after L1 and the university's masters."""
    if fili not in SU.LICENCE_FILI:
        return None
    has_reo, has_mm = _ready(con, "core.licence_reorientation"), _ready(con, "core.monmaster")
    if not (has_reo or has_mm):
        return None
    try:
        link = con.execute("""SELECT etab, discipline, lien FROM core.licence_success_link
                              WHERE cod_aff_form = ? ORDER BY session DESC LIMIT 1""", [cod]).fetchone()
    except Exception:  # noqa: BLE001
        link = None
    etab, disc, lien = link if link else (None, None, None)
    out = {"status": "ok", "etab": etab, "etab_label": SU.UNIVERSITY_LABEL.get(etab, etab) if etab else None,
           "discipline": disc, "discipline_label": SU.DISCIPLINES[disc][1] if disc else None, "lien": lien,
           "reorientation": None, "reorientation_france": None, "masters": None, "masters_france": None}
    if contrat != "Public":
        out.update(status="hors_champ", reason="le SIES publie ces chiffres pour les universités publiques seulement")
        return out
    if etab is None:
        out.update(status="non_rattachee", reason="cet établissement n'a pas pu être rattaché avec certitude à une université des tableaux SIES")
        return out
    if disc is None:
        out.update(status="discipline_non_classee", reason="cette licence ne relève pas clairement d'une seule des cinq grandes disciplines du SIES")
        return out
    if has_reo:
        out["reorientation"] = _reo_row(con, etab, disc)
        out["reorientation_france"] = _reo_row(con, NATIONAL, disc)
        r = out["reorientation"]
        if r:
            r["maillon"] = E.maillon("reorientation_l1", r["taux_reorientation"], france=(out["reorientation_france"] or {}).get("taux_reorientation"),
                                     granularite="universite_x_discipline", population="néo-bacheliers inscrits en L1",
                                     cohorte=r["cohort"], effectif=r["inscrits"], source="SIES, Note Flash 2025-29", lien=lien or "nom",
                                     avertissement="tous les inscrits en L1 de cette grande discipline dans l'université")
    if has_mm:
        s = latest_session(con)
        pid = (SU.SIES_PAYSAGE if paysage is None else paysage).get(etab)
        m = masters_of(con, pid, disc, s) if pid else None
        out["masters"], out["masters_france"] = m, masters_of(con, None, disc, s)
        out["meef"] = masters_of(con, pid, "meef", s) if pid else None
        if m:
            m["maillon"] = E.maillon("admis_master", m["part_etab"], france=(out["masters_france"] or {}).get("part_etab"),
                                     granularite="universite_x_discipline", population="admis en master (MonMaster)",
                                     cohorte=m["session"], effectif=m["n_accept"], source="SIES, MonMaster", lien=lien or "nom",
                                     avertissement="tous les masters de cette grande discipline dans l'université")
    return out


def for_but(con, etab: str | None, paysage: dict[str, str] | None = None) -> dict | None:
    """BUT 3 graduates admitted in the masters of the BUT's university (and in France), latest session."""
    if not _ready(con, "core.monmaster"):
        return None
    s = latest_session(con)
    nat = con.execute("SELECT sum(n_accept), sum(n_accept_but3) FROM core.monmaster WHERE session = ?", [s]).fetchone()
    out = {"session": int(s), "france": {"n_accept": nat[0], "n_accept_but3": nat[1],
                                         "part_but3": 100 * nat[1] / nat[0] if nat[0] else None}, "universite": None}
    pid = (SU.SIES_PAYSAGE if paysage is None else paysage).get(etab) if etab else None
    if pid:
        u = con.execute("SELECT sum(n_accept), sum(n_accept_but3) FROM core.monmaster WHERE session = ? AND id_paysage = ?", [s, pid]).fetchone()
        if u[0]:
            out["universite"] = {"n_accept": u[0], "n_accept_but3": u[1], "part_but3": 100 * u[1] / u[0]}
    return out
