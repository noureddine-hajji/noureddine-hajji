"""SYNTHETIC raw files that mimic the SIES export layout, for tests only.

Nothing here is real Parcoursup data. The generator deliberately injects the
situations the pipeline must survive: column-name drift between vintages, a
re-coded cod_aff_form, an ambiguous re-coding, a campaign gap, a concours commun
sharing one applicant pool, and low-denominator cohorts.
"""
from __future__ import annotations

import random
from pathlib import Path

import pandas as pd

YEARS = (2022, 2023, 2024, 2025)

ETABS = [
    ("0751234A", "Université Exemple Paris", "Paris", "75", "Paris"),
    ("0691234B", "IUT Exemple Lyon", "Lyon", "69", "Lyon"),
    ("0331234C", "Lycée Exemple Bordeaux", "Bordeaux", "33", "Bordeaux"),
    ("0591234D", "École d'ingénieurs Exemple Lille", "Lille", "59", "Lille"),
    ("0131234E", "Lycée Exemple Marseille", "Marseille", "13", "Aix-Marseille"),
    ("0311234F", "École d'ingénieurs Exemple Toulouse", "Toulouse", "31", "Toulouse"),
    ("0441234G", "École d'ingénieurs Exemple Nantes", "Nantes", "44", "Nantes"),
    ("0671234H", "École d'ingénieurs Exemple Strasbourg", "Strasbourg", "67", "Strasbourg"),
]

TRACKS = [  # fili, form_lib_voe_acc, fil_lib_voe_acc (real SIES taxonomy labels)
    ("BUT", "BUT - Production", "Génie mécanique et productique"),
    ("BUT", "BUT - Service", "Gestion des entreprises et des administrations"),
    ("BTS", "BTS - Services", "Comptabilité et gestion"),
    ("BTS", "BTS - Production", "Systèmes numériques"),
    ("Licence", "Licence - Droit-économie-gestion", "Droit"),
    ("Licence", "Licence - Sciences - technologies - santé", "Mathématiques"),
    ("CPGE", "Classe préparatoire scientifique", "MPSI"),
]
ECG = ("CPGE", "Classe préparatoire économique et commerciale", "ECG - Mathématiques appliquées + ESH")
ECT = ("CPGE", "Classe préparatoire économique et commerciale", "ECT - Option technologique")
AES = ("Licence", "Licence - Droit-économie-gestion", "Administration économique et sociale")
DOUBLETTES = {  # SIES spelling
    "MATHS+SES": "Mathématiques Spécialité,Sciences Economiques et Sociales Spécialité",
    "MATHS+PC": "Mathématiques Spécialité,Physique-Chimie Spécialité",
    "HGGSP+SES": "Sciences Economiques et Sociales Spécialité,Histoire-Géographie, Géopolitique et Sciences politiques",
}
# (voeux, propositions, acceptations) per year, by SIES formation type and doublette key
SPE_PLAN = {
    "CPGE ECG": {"MATHS+SES": (900, 700, 600), "MATHS+PC": (300, 150, 100), "HGGSP+SES": (100, 40, 20)},
    "CPGE S": {"MATHS+SES": (40, 3, 2), "MATHS+PC": (3000, 2400, 2000), "HGGSP+SES": (5, 0, 0)},
    "Licence Pluri Sciences économiques - Gestion":
        {"MATHS+SES": (1500, 1200, 800), "MATHS+PC": (300, 250, 150), "HGGSP+SES": (900, 700, 500)},
    "Licence AES": {"MATHS+SES": (400, 320, 200), "MATHS+PC": (100, 80, 40), "HGGSP+SES": (600, 450, 300)},
    "Licence Sciences juridiques":
        {"MATHS+SES": (500, 350, 250), "MATHS+PC": (500, 300, 200), "HGGSP+SES": (3000, 2000, 1500)},
    "BUT Service": {"MATHS+SES": (800, 450, 300), "MATHS+PC": (300, 150, 80), "HGGSP+SES": (500, 250, 150)},
}
NATIONAL = {"MATHS+SES": (3000, 2900), "MATHS+PC": (3600, 3500), "HGGSP+SES": (3500, 3400)}  # voeux, prop
SPLIT_LABEL = "Bac Série générale avec la spécialité Maths + une spécialité non scientifique"


def _fixed_cohort(sm, ab, b, tb, tbf, bg, bt, bp):
    """Admitted cohort with an exact mention / series profile (identities hold)."""
    neo = sm + ab + b + tb + tbf
    assert bg + bt + bp == neo
    pct = lambda x: round(100.0 * x / neo) if neo else None  # noqa: E731
    return {"acc_tot": neo, "acc_tot_f": neo // 2, "acc_term": neo, "acc_term_f": neo // 2, "acc_neobac": neo,
            "acc_at": 0, "acc_bg": bg, "acc_bt": bt, "acc_bp": bp, "acc_brs": neo // 4, "acc_pp": neo,
            "acc_pc": 0, "acc_aca_orig": neo // 2, "acc_internat": 0, "acc_sansmention": sm, "acc_ab": ab,
            "acc_b": b, "acc_tb": tb, "acc_tbf": tbf, "acc_mention_nonrenseignee": 0,
            "pct_f": 50, "pct_bours": 25, "pct_bg": pct(bg), "pct_bt": pct(bt), "pct_bp": pct(bp),
            "pct_tb": pct(tb), "pct_tbf": pct(tbf), "pct_b": pct(b), "pct_ab": pct(ab), "pct_sansmention": pct(sm),
            "pct_mention_nonrenseignee": 0, "pct_aca_orig": 50, "pct_etab_orig": 10, "pct_neobac": 100}


def _series_mentions(rng: random.Random, c: dict) -> dict:
    """Split the admits with a mention (AB..TBF) across series, within each series' size."""
    m = sum(c.get(k, 0) for k in ("acc_ab", "acc_b", "acc_tb", "acc_tbf"))
    bg = min(c["acc_bg"], rng.randint(max(0, m - c["acc_bt"] - c["acc_bp"]), m))
    bt = min(c["acc_bt"], m - bg)
    bp = m - bg - bt
    return {"acc_bg_mention": bg, "acc_bt_mention": bt, "acc_bp_mention": bp,
            # like SIES: over all néo-bacheliers admitted, not over the bac-général admits
            "pct_bg_mention": round(100.0 * bg / c["acc_neobac"]) if c["acc_neobac"] else None}


def _series_candidates(rng: random.Random, r: dict) -> dict:
    """Candidates (PP, PC) and proposals by series, consistent like the real exports: admits <=
    proposals <= candidates for each series, and the series fit inside the formation totals
    (which stay untouched: a shared concours pool must keep identical totals)."""
    ser = ("bg", "bt", "bp")
    acc = {s: r.get(f"acc_{s}", 0) or 0 for s in ser}
    left = {"pp": r["nb_voe_pp"] - sum(acc.values()), "prop": r["prop_tot"] - sum(acc.values()),
            "pc": r["nb_voe_pc"]}
    out = {}
    for s in ser:                       # every admit was a main-phase candidate who got a proposal
        extra_pp = rng.randint(0, max(0, left["pp"]) // 3)
        extra_prop = min(rng.randint(0, max(0, left["prop"]) // 3), acc[s] + extra_pp - acc[s])
        c_pc = rng.randint(0, max(0, left["pc"]) // 3)
        out.update({f"nb_voe_pp_{s}": acc[s] + extra_pp, f"nb_voe_pc_{s}": c_pc, f"prop_tot_{s}": acc[s] + extra_prop})
        left["pp"] -= extra_pp
        left["prop"] -= extra_prop
        left["pc"] -= c_pc
    return out


def _timing(rng: random.Random, acc_tot: int) -> dict:
    """Cumulative offer timing, as published: counts and SIES-rounded shares of acc_tot."""
    d = rng.randint(0, acc_tot)
    b = rng.randint(d, acc_tot)
    f = rng.randint(b, acc_tot)
    pct = lambda x: round(100.0 * x / acc_tot) if acc_tot else None  # noqa: E731
    return {"acc_debutpp": d, "acc_datebac": b, "acc_finpp": f,
            "pct_acc_debutpp": pct(d), "pct_acc_datebac": pct(b), "pct_acc_finpp": pct(f),
            "nb_voe_pp_internat": 0, "nb_cla_pp_internat": 0}


def _cohort(rng: random.Random, acc_tot: int) -> dict:
    acc_at = rng.randint(0, max(0, acc_tot // 6))
    term = acc_tot - acc_at
    bg = rng.randint(term // 3, term)
    bt = rng.randint(0, term - bg)
    bp = term - bg - bt
    rest = term
    mentions = {}
    for k in ("acc_tbf", "acc_tb", "acc_b", "acc_ab", "acc_mention_nonrenseignee"):
        v = rng.randint(0, rest // 3)
        mentions[k] = v
        rest -= v
    mentions["acc_sansmention"] = rest
    brs = rng.randint(0, term)
    f = rng.randint(0, acc_tot)
    pct = lambda x, d: round(100.0 * x / d) if d else None  # noqa: E731 - SIES rounds to integers
    return {
        # like the real SIES exports: breakdowns add up to acc_neobac, acc_term does not
        "acc_tot": acc_tot, "acc_tot_f": f, "acc_term": max(0, term - rng.randint(0, 3)),
        "acc_term_f": min(f, term), "acc_neobac": term, "acc_at": acc_at, "acc_bg": bg, "acc_bt": bt, "acc_bp": bp,
        "acc_brs": brs, "acc_pp": acc_tot - acc_tot // 10, "acc_pc": acc_tot // 10,
        "acc_aca_orig": rng.randint(0, acc_tot), "acc_internat": 0, **mentions, **_timing(rng, acc_tot),
        "pct_f": pct(f, acc_tot), "pct_bours": pct(brs, term), "pct_bg": pct(bg, term),
        "pct_bt": pct(bt, term), "pct_bp": pct(bp, term),
        "pct_tb": pct(mentions["acc_tb"], term), "pct_tbf": pct(mentions["acc_tbf"], term),
        "pct_b": pct(mentions["acc_b"], term), "pct_ab": pct(mentions["acc_ab"], term),
        "pct_sansmention": pct(mentions["acc_sansmention"], term),
        "pct_mention_nonrenseignee": pct(mentions["acc_mention_nonrenseignee"], term),
        "pct_aca_orig": rng.randint(20, 95), "pct_etab_orig": rng.randint(0, 40),
        "pct_neobac": pct(term, acc_tot),
    }


def _row(rng, year, cod, etab, track, lib_comp="", acc_tot=None, pool=None):
    uai, name, ville, dep, aca = etab
    fili, form, fil = track
    acc_tot = acc_tot if acc_tot is not None else rng.randint(12, 120)
    capa = acc_tot + rng.randint(0, 10)
    if pool:
        voe_tot, voe_tot_f, nb_pp, nb_pc = pool
    else:
        nb_pp = rng.randint(capa * 2, capa * 30)
        nb_pc = rng.randint(0, capa)
        voe_tot, voe_tot_f = nb_pp + nb_pc, rng.randint(0, nb_pp + nb_pc)
    nb_cla = rng.randint(capa, nb_pp)
    row = {
        "session": year, "cod_uai": uai, "cod_aff_form": str(cod), "g_ea_lib_vx": name,
        "contrat_etab": "Public", "dep": dep, "dep_lib": ville, "region_etab_aff": "Région",
        "acad_mies": aca, "ville_etab": ville, "select_form": "formation sélective",
        "fili": fili, "form_lib_voe_acc": form, "fil_lib_voe_acc": fil,
        "lib_for_voe_ins": f"{fil} {lib_comp}".strip(), "lib_comp_voe_ins": lib_comp,
        "detail_forma": None, "capa_fin": capa, "voe_tot": voe_tot, "voe_tot_f": voe_tot_f,
        "nb_voe_pp": nb_pp, "nb_voe_pc": nb_pc, "nb_cla_pp": nb_cla, "nb_cla_pc": 0,
        "prop_tot": min(nb_cla, capa * 3), "ran_grp1": rng.randint(capa, nb_cla),
        "taux_acces_ens": round(rng.uniform(5, 100)), "part_acces_gen": round(rng.uniform(0, 100)),
        "part_acces_tec": round(rng.uniform(0, 100)), "part_acces_pro": round(rng.uniform(0, 100)),
        "g_olocalisation_des_formations": "48.85, 2.35",
        "lien_form_psup": ("https://dossierappel.parcoursup.fr/Candidats/public/fiches/"
                           f"afficherFicheFormation?g_ta_cod={cod}&typeBac=0&originePc=0"),
        **_cohort(rng, acc_tot),
    }
    row.update(_series_mentions(rng, row))
    row.update(_series_candidates(rng, row))
    return row


def make_raw(raw_dir: Path, seed: int = 7) -> dict:
    """Write synthetic raw parquet files; returns the planted scenario for assertions."""
    rng = random.Random(seed)
    raw_dir.mkdir(parents=True, exist_ok=True)
    plan = {"recoded": ("1003", "9003"), "ambiguous": ("1010", "1011"), "gap": "1020",
            "new_2025": "1030", "low_n": "1040", "pool": ["2001", "2002", "2003", "2004"],
            "ecg_elite": "3001", "ecg_open": "3002", "ect": "3003", "ecg_mid": "3006", "aes": "3005",
            "ecogestion": "3007", "pcsi": "3008"}
    stable = []
    cod = 1100
    for e in ETABS[:5]:
        for t in TRACKS:
            stable.append((str(cod), e, t))
            cod += 1
    pools = {y: (900 + y % 100, 400, 880 + y % 100, 20) for y in YEARS}
    for y in YEARS:
        rows = [_row(rng, y, c, e, t) for c, e, t in stable]
        # re-coded in 2024: same UAI + fili + labels, new cod_aff_form
        rows.append(_row(rng, y, "1003" if y < 2024 else "9003", ETABS[0], TRACKS[4],
                         lib_comp="Parcours droit public"))
        # two programmes with identical labels both re-coded in 2024 -> must NOT be linked
        for old, new in (("1010", "9010"), ("1011", "9011")):
            rows.append(_row(rng, y, old if y < 2024 else new, ETABS[1], TRACKS[1]))
        if y != 2023:
            rows.append(_row(rng, y, "1020", ETABS[2], TRACKS[2], lib_comp="Alternance"))
        if y == 2025:
            rows.append(_row(rng, y, "1030", ETABS[3], TRACKS[5]))
        rows.append(_row(rng, y, "1040", ETABS[4], TRACKS[6], acc_tot=6))
        for c, e in zip(plan["pool"], ETABS[3:]):
            rows.append(_row(rng, y, c, e, ("Ecole d'Ingénieur", "Formation d'ingénieur Bac + 5",
                                            "Formation d'ingénieur Bac + 5 - bac général"),
                             lib_comp="Concours Puissance Alpha", pool=pools[y]))
        # journey personas: an elite and an accessible ECG, a mid one, an ECT and an AES licence
        journey = [
            ("3001", ETABS[0], ECG, 5, _fixed_cohort(0, 0, 5, 25, 15, 45, 0, 0), None),
            ("3002", ETABS[1], ECG, 85, _fixed_cohort(15, 20, 8, 2, 0, 40, 5, 0),
             [(SPLIT_LABEL, 120), ("Autres candidats", 300)]),
            ("3006", ETABS[2], ECG, 45, _fixed_cohort(4, 10, 16, 12, 3, 45, 0, 0), None),
            ("3009", ETABS[3], ECG, 20, _fixed_cohort(1, 4, 15, 18, 7, 45, 0, 0), None),
            ("3003", ETABS[3], ECT, 60, _fixed_cohort(15, 15, 8, 2, 0, 0, 40, 0), None),
            ("3005", ETABS[4], AES, 90, _fixed_cohort(60, 50, 20, 5, 1, 120, 12, 4), None),
            ("3007", ETABS[0], ("Licence", "Licence - Droit-économie-gestion", "Economie et gestion"), 80,
             _fixed_cohort(50, 60, 30, 8, 2, 140, 10, 0), None),
            ("3008", ETABS[1], ("CPGE", "Classe préparatoire scientifique", "PCSI"), 40,
             _fixed_cohort(0, 2, 10, 20, 10, 42, 0, 0), None),
        ]
        for c, e, track, taux, cohort, groups in journey:
            r = _row(rng, y, c, e, track)
            r.update(cohort)
            r.update(_timing(rng, cohort["acc_tot"]))
            r.update(_series_mentions(rng, r))
            # a fixed cohort can exceed the random row's proposals and candidates: admits never do
            r["prop_tot"] = max(r["prop_tot"], r["acc_tot"])
            r["nb_voe_pp"] = max(r["nb_voe_pp"], r["acc_tot"])
            r["voe_tot"] = max(r["voe_tot"], r["nb_voe_pp"] + r["nb_voe_pc"])
            r.update(_series_candidates(rng, r))
            r.update({"taux_acces_ens": taux, "capa_fin": cohort["acc_tot"] + 2,
                      "lib_grp1": "Tous les candidats", "ran_grp1": 400, "lib_grp2": None, "ran_grp2": None})
            for i, (lib, ran) in enumerate(groups or [], start=1):
                r[f"lib_grp{i}"], r[f"ran_grp{i}"] = lib, ran
            rows.append(r)
        df = pd.DataFrame(rows)
        # Paysage ids as SIES publishes them: filled 2022-2024, EMPTY in 2025; Paris changes id in 2023
        # (a renaming seen as a transition), the Lyon IUT keeps a former id that only its UAI resolves
        # (the register lists that UAI as former UAI of LYON1); lycées carry none.
        # Two traps for the succession rules: "SPLIT" is carried by two lycées whose UAIs the register
        # gives (as former UAIs) to two different establishments, half and half (no succession); "TINY1"
        # is carried by a single formation (too little evidence).
        ids = {"0751234A": "OLDP1" if y == 2022 else "PARIS", "0691234B": "LYONX", "0591234D": "LILLE",
               "0331234C": "SPLIT", "0131234E": "SPLIT", "0311234F": "TINY1"}
        df["etablissement_id_paysage"] = None if y == 2025 else df.cod_uai.map(ids)
        df["composante_id_paysage"] = None
        # synthetic schema drift between vintages
        if y == 2022:
            df.columns = [c.upper() for c in df.columns]
        if y == 2023:
            df = df.rename(columns={"acc_mention_nonrenseignee": "acc_mention_non_renseignee",
                                    "taux_acces_ens": "taux_acces"})
        if y == 2025:
            df = df.rename(columns={"part_acces_gen": "part_terminale_generale_en_position_admis_pp"})
        df.to_parquet(raw_dir / f"parcoursup_{y}.parquet", index=False)
        if y == 2025:  # real exports carry GEOMETRY('OGC:CRS84') points
            import duckdb
            path = raw_dir / f"parcoursup_{y}.parquet"
            duckdb.sql(f"""COPY (SELECT * REPLACE ('POINT (2.35 48.85)'::GEOMETRY AS g_olocalisation_des_formations)
                           FROM read_parquet('{path}')) TO '{path}.tmp' (FORMAT parquet)""")
            Path(f"{path}.tmp").replace(path)

    spe = []
    for y in YEARS:
        for ft, by_d in SPE_PLAN.items():
            for dk, (v, pr, a) in by_d.items():
                spe.append({"annee_du_bac": y, "formation": ft, "doublette": DOUBLETTES[dk],
                            "voeux": v, "propositions_d_admissions": pr, "acceptations": a})
        for dk, (v, pr) in NATIONAL.items():   # deduplicated national row: each candidate accepts once
            spe.append({"annee_du_bac": y, "formation": "Ensemble des bacheliers", "doublette": DOUBLETTES[dk],
                        "voeux": v, "propositions_d_admissions": pr,
                        "acceptations": sum(SPE_PLAN[ft][dk][2] for ft in SPE_PLAN)})
    spe = pd.DataFrame(spe)
    spe[spe.annee_du_bac < 2025].to_parquet(raw_dir / "specialites_2021_2024.parquet", index=False)
    # 2025 layout: level 0 = all bacheliers, 1 = group total, 2 = detailed formations (repeat level 1)
    s25 = spe[spe.annee_du_bac == 2025].rename(columns={"formation": "regroupement_de_formations"})
    nat = s25.regroupement_de_formations == "Ensemble des bacheliers"
    lvl0 = s25[nat].assign(niveau_d_agregation="0", formation="Ensemble des bacheliers")
    lvl1 = s25[~nat].assign(niveau_d_agregation="1", formation="Ensemble regroupement de formations")
    lvl2 = s25[~nat].assign(niveau_d_agregation="2", formation=s25.regroupement_de_formations + " (détail)")
    pd.concat([lvl0, lvl1, lvl2]).to_parquet(raw_dir / "specialites_2025.parquet", index=False)

    fiche = ("https://dossierappel.parcoursup.fr/Candidats/public/fiches/"
             "afficherFicheFormation?g_ta_cod={}&typeBac=0&originePc=0")
    # real ODS values; the cartographie lists several sessions per formation (latest wins)
    carto = [
        {"gta": "1100", "annee": 2025, "etab_gps": "45.75, 4.85", "fiche": fiche.format(1100),
         "app": "", "int": "", "amg": "Enseignement en présentiel"},
        {"gta": "1100", "annee": 2026, "etab_gps": "45.75, 4.85", "fiche": fiche.format(1100),
         "app": "", "int": "Etablissements avec internat pour filles et garçons",
         "amg": "Enseignement en présentiel,Formations avec aménagements pour sportifs haut niveau"},
        {"gta": "1102", "annee": 2026, "etab_gps": None, "fiche": fiche.format(1102),
         "app": "Formations en apprentissage", "int": "",
         "amg": "Enseignement partiellement à distance,Formations avec aménagements pour artistes haut niveau"},
        {"gta": "1103", "annee": 2026, "etab_gps": None, "fiche": fiche.format(1103),
         "app": "", "int": "Etablissements avec internat pour filles", "amg": "Enseignement à distance"},
    ]
    for c in carto:   # the cartographie carries the current Paysage id (fills the empty 2025 export)
        c["etablissement_id_paysage"] = "PARIS" if c["gta"] in ("1100", "1102") else None
        c["composante_id_paysage"] = None
    pd.DataFrame(carto).to_parquet(raw_dir / "cartographie.parquet", index=False)
    register = [  # MESR "Principaux établissements" layout (subset of columns)
        {"etablissement_id_paysage": "PARIS", "uai": "0751234A", "anciens_codes_uai": None, "uo_lib": "Université Exemple Paris",
         "type_d_etablissement": "Université", "secteur_d_etablissement": "public", "aca_nom": "Paris"},
        {"etablissement_id_paysage": "LYON1", "uai": "0699999Z", "anciens_codes_uai": "0691234B;0690000X",
         "uo_lib": "Université Exemple Lyon", "type_d_etablissement": "Université", "secteur_d_etablissement": "public", "aca_nom": "Lyon"},
        {"etablissement_id_paysage": "LILLE", "uai": "0591234D", "anciens_codes_uai": None, "uo_lib": "École d'ingénieurs Exemple Lille",
         "type_d_etablissement": "École", "secteur_d_etablissement": "public", "aca_nom": "Lille"},
        {"etablissement_id_paysage": "BORD1", "uai": "0339999Z", "anciens_codes_uai": "0331234C", "uo_lib": "Université Exemple Bordeaux",
         "type_d_etablissement": "Université", "secteur_d_etablissement": "public", "aca_nom": "Bordeaux"},
        {"etablissement_id_paysage": "MARS1", "uai": "0139999Z", "anciens_codes_uai": "0131234E", "uo_lib": "Université Exemple Marseille",
         "type_d_etablissement": "Université", "secteur_d_etablissement": "public", "aca_nom": "Aix-Marseille"},
        {"etablissement_id_paysage": "TOUL1", "uai": "0319999Z", "anciens_codes_uai": "0311234F", "uo_lib": "Université Exemple Toulouse",
         "type_d_etablissement": "Université", "secteur_d_etablissement": "public", "aca_nom": "Toulouse"},
        {"etablissement_id_paysage": "PRIVE", "uai": "0759999Z", "anciens_codes_uai": None, "uo_lib": "École privée Exemple",
         "type_d_etablissement": "École", "secteur_d_etablissement": "privé", "aca_nom": "Paris"},
    ]
    pd.DataFrame(register).to_parquet(raw_dir / "etablissements.parquet", index=False)
    plan["paysage"] = {"transition": ("OLDP1", "PARIS"), "uai": ("LYONX", "LYON1"), "carto": ("1102", "PARIS")}
    # Sprint 4 lycée files, real column names and quirks: three school years, numbers as text, an LP
    # without GT-track IPS, a private lycée, an académie with no Parcoursup offer; IVAL with a
    # truncated académie label, old sessions and a negative value added. The Bordeaux lycée hosts
    # synthetic BTS formations (same-lycée share).
    ips = []
    for rentree in ("2023-2024", "2024-2025", "2025-2026"):
        for uai, nom, com, aca, typ, sect, gt, et in (
                ("0331234C", "LYCEE EXEMPLE BORDEAUX", "BORDEAUX", "BORDEAUX", "LPO", "public", "112.4", "110.1"),
                ("0131234E", "LYCEE EXEMPLE MARSEILLE", "MARSEILLE", "AIX-MARSEILLE", "LEGT", "privé sous contrat", "131.0", "131.0"),
                ("0339876P", "LYCEE PROFESSIONNEL EXEMPLE", "MERIGNAC", "BORDEAUX", "LP", "public", None, "84.2"),
                ("9750001A", "LYCEE EXEMPLE SAINT-PIERRE", "SAINT-PIERRE", "ST PIERRE ET MIQUELON", "LPO", "public", "99.0", "98.0")):
            if uai == "0339876P" and rentree == "2023-2024":
                continue
            row = {"rentree_scolaire": rentree, "academie": aca, "nom_de_la_commune": com, "uai": uai,
                   "nom_de_l_etablissement": nom, "secteur": sect, "type_de_lycee": typ, "ips_voie_gt": gt,
                   "ips_voie_pro": "90.0" if typ != "LEGT" else None, "ips_etab": et, "ecart_type_etablissement": "30.1"}
            for lvl, base in (("national", 100.0), ("academique", 101.0)):
                for t in ("legt", "lpo", "lp"):
                    for sk, bonus in (("public", 0.0), ("prive", 15.0)):
                        row[f"ips_{lvl}_{t}_{sk}"] = base + {"legt": 15, "lpo": 3, "lp": -12}[t] + bonus
            ips.append(row)
    pd.DataFrame(ips).to_parquet(raw_dir / "ips_lycees.parquet", index=False)
    ival = []
    for year in (2012, 2023, 2024, 2025):
        for uai, aca, va in (("0331234C", "BORDEAUX", -3.0), ("0131234E", "AIX-MARSEILLE", 4.0), ("0130000Z", "AIX-MARSEILLE", 1.0)):
            if uai == "0130000Z" and year != 2012:      # a lycée closed long ago
                continue
            ival.append({"annee": f"{year}-01-01", "uai": uai, "libelle_academie": "NORMAIE" if uai == "0331234C" else aca,
                         "presents_total": 25 if uai == "0131234E" else 60 if year == 2023 else 250, "taux_reu_total": 92.0, "va_reu_total": va,
                         "taux_men_total": 60.0, "va_men_total": va / 2, "taux_acces_2nde": 85.0, "va_acces_2nde": -1.0})
    df = pd.DataFrame(ival)
    df["annee"] = pd.to_datetime(df["annee"]).dt.date
    df.to_parquet(raw_dir / "ival_lycees_gt.parquet", index=False)
    return plan


# ---------------------------------------------------------------- B1: SIES licence success workbook
SUCCESS_TABLE = {"UNIVERSITE EXEMPLE PARIS (CITE)": r"^universite exemple paris", "UNIVERSITE EXEMPLE NORD": r"exemple nord"}


def make_success_xlsx(path: Path) -> dict:
    """A workbook with the real file's layout and quirks: title + 'Retour au sommaire' + header rows,
    the two national blocks ("Etudiants", then "Inscriptions principales" = sum of universities),
    a university whose discipline rows carry a shorter name than its total row, masked cells
    ("ns", "nd", "."), a newline inside a discipline label, and footers."""
    import openpyxl
    unis = {  # name of the total row, name of the discipline rows, {discipline: entrants}
        "UNIVERSITE EXEMPLE PARIS (CITE)": ("UNIVERSITE EXEMPLE PARIS", {"Droit, Sc. politiques": 400, "Sciences-santé": 300, "STAPS": 60}),
        "UNIVERSITE EXEMPLE NORD": ("UNIVERSITE EXEMPLE NORD", {"Droit, Sc. politiques": 200, "Arts, lettres, langues, SHS": 90}),
    }
    rates = {"Droit, Sc. politiques": (50.0, 45.0, 25.0, 12.0), "Sciences-santé": (40.0, 44.0, 30.0, 15.0),
             "STAPS": (55.0, 50.0, 20.0, 10.0), "Arts, lettres, langues, SHS": (45.0, 47.0, 28.0, 11.0)}
    wb = openpyxl.Workbook()
    wb.active.title = "Sommaire"
    wb.active.append(["Tableaux par discipline et établissement - Sommaire"])
    plan = {"rows": {}}

    def block(ws, kind, total_name, short, discs, national=False):
        tot = {"n": 0}
        lines = []
        for d, n in discs.items():
            obs, sim, red, reo = rates[d]
            masked = kind != "passage" and n < 100           # SIES masks success below ~100 entrants
            if kind == "passage":
                p, r_, o = round(n * obs / 100), round(n * red / 100), round(n * reo / 100)
                s_ = n - p - r_ - o
                va = "." if national else round(round(100 * p / n, 1) - sim, 1)
                sim_v = "." if national else sim
                line = [short, d.replace(", ", ",\n", 1) if d.startswith("Droit") else d, n, p, round(100 * p / n, 1), sim_v, va,
                        p, round(100 * p / n, 1), 0, 0.0, r_, round(100 * r_ / n, 1), r_, round(100 * r_ / n, 1),
                        o, round(100 * o / n, 1), s_, round(100 * s_ / n, 1)]
            elif kind == "reussite3":
                k = round(n * (obs - 10) / 100)
                line = [short, d, n] + (["ns"] * 5 if masked else [k, round(100 * k / n, 1), "." if national else sim - 10,
                                                                   "." if national else round(round(100 * k / n, 1) - (sim - 10), 1), 0.5])
            else:
                k3, k4 = round(n * (obs - 12) / 100), round(n * 8 / 100)
                line = [short, d, n] + (["ns"] * 11 if masked else [
                    k3, round(100 * k3 / n, 1), "." if national else sim - 12, "." if national else round(round(100 * k3 / n, 1) - (sim - 12), 1), 0.4,
                    k4, round(100 * k4 / n, 1), k3 + k4, round(100 * (k3 + k4) / n, 1), "." if national else sim - 4,
                    "." if national else round(round(100 * (k3 + k4) / n, 1) - (sim - 4), 1)])
            lines.append(line)
            tot["n"] += n
            plan["rows"][(kind, total_name, d)] = line
        n = tot["n"]
        total = [total_name, "Toutes disciplines", n] + ["nd"] * (len(lines[0]) - 3)
        ws.append(total)
        for line in lines:
            ws.append(line)
        ws.append([])

    sheets = {
        "passage": ("Devenir cohorte 2023 VA", "Devenir en 2024-2025 des néo-bacheliers inscrits en licence à la rentrée 2023 selon établissement et discipline en L1",
                    ["Nom établissement en L1", "Discipline en L1", "Inscrits en L1", "Passage en L2", "Taux de passage en L2", "Taux de passage en L2 simulé",
                     "Passage en L2 - Valeur Ajoutée (en points)", "Dont passage strict", "Taux de passage strict", "Dont passage en MMOPK (LAS)",
                     "Taux de passage en MMOPK", "Redoublement en L1", "Taux de redoublement en L1", "Dont redoublement strict", "Taux de redoublement strict",
                     "Réorientation", "Taux de réorientation", "Sortie de l'enseignement supérieur", "Taux de sortie de l'enseignement supérieur"]),
        "reussite3": ("Réussite 3ans discip VA", "Réussite en licence en 3 ans des néo-bacheliers inscrits en licence en 2021 qui n'ont pas changé de discipline - corrigé",
                      ["Nom établissement en L1", "Discipline en L1", "Inscrits en L1 ", "Réussite en 3 ans", "Taux de réussite à 3 ans", "Taux simulé",
                       "Valeur ajoutée (en points)", "Passage en MMOPK (%)"]),
        "reussite34": ("Réussite 3ou4ans discip VA", "Réussite en licence en 3 ou 4 ans des néo-bacheliers inscrits en licence en 2020 qui n'ont pas changé de discipline",
                       ["Nom établissement en L1", "Discipline en L1", "Inscrits en L1", "Réussite en 3 ans", "Réussite en 3 ans (taux observé)",
                        "Réussite en 3 ans \n(taux simulé)", "Valeur ajoutée réussite en 3 ans (en points)", "Passage en MMOPK en 3 ans (%)", "Réussite en 4 ans",
                        "Réussite en 4 ans (taux observé)", "Réussite en 3 ou 4 ans", "Réussite en 3 ou 4 ans (taux observé)", "Réussite en 3 ou 4 ans \n(taux simulé)",
                        "Valeur ajoutée réussite en 3 ou 4 ans (en points)"]),
    }
    alld = {}
    for _, (_, ds) in unis.items():
        for d, n in ds.items():
            alld[d] = alld.get(d, 0) + n
    for kind, (title, long_title, header) in sheets.items():
        ws = wb.create_sheet(title)
        ws.append([long_title])
        ws.append(["Retour au sommaire"])
        ws.append(header)
        block(ws, kind, "Etudiants", "Etudiants", {d: n - 5 for d, n in alld.items()}, national=True)
        block(ws, kind, "Inscriptions principales", "Inscriptions principales", alld, national=True)
        for name, (short, ds) in unis.items():
            block(ws, kind, name, short, ds)
        ws.append(["ns : résultat non significatif du fait de faibles effectifs"])
        ws.append(["Source : MESRE-SIES"])
    # sprint 7: where the reoriented went (same entrants and reorientations as the "Devenir" sheet; the
    # Nord university's total row carries another spelling, as Paris 13 does in the real sheet)
    ws = wb.create_sheet("Réorientations - cohorte 2023")
    ws.append(["Détail des formations suivies en 2024-2025 par les néo-bacheliers inscrits en licence à la rentrée 2023 et réorientés à la rentrée 2024 selon établissement et discipline en L1 - corrigé"])
    ws.append(["Retour au sommaire"])
    ws.append(["Nom établissement en L1", "Discipline en L1", "Inscrits en L1\n(rappel)", "Réorientation\n(rappel)", "Taux de réorientation\n(rappel)",
               "Dont réorientation hors université", "Taux de réorientation hors université", "Dont réorientation en CPGE", "Taux de réorientation en CPGE",
               "Dont réorientation en STS en apprentissage", "Taux de réorientation en STS en apprentissage", "Dont réorientation en STS hors apprentissage",
               "Taux de réorientation en STS hors apprentissage", "Dont réorientation en école de commerce", "Taux de réorientation en école de commerce",
               "Dont réorientation en école d'ingénieur", "Taux de réorientation en école d'ingénieur", "Dont réorientation à l'université",
               "Taux de réorientation à l'université", "Dont réorientation en DUT/BUT", "Taux de réorientation en DUT/BUT",
               "Dont réorientation en DE Infirmier", "Taux de réorientation en DE Infirmier"])

    def reo_line(name, d, n):
        o = round(n * rates[d][3] / 100)
        if n < 100:
            return [name, d, n] + ["ns"] * 20
        hors = round(o * 0.6)
        univ = o - hors
        parts = [hors, 0, round(hors * 0.3), round(hors * 0.3), round(hors * 0.1), 0, univ, round(univ * 0.4), round(univ * 0.2)]
        line = [name, d, n, o, round(100 * o / n, 1)]
        for x in parts:
            line += [x, round(100 * x / n, 1)]
        return line

    plan["reo"] = {}
    for total, (short, ds) in [("Etudiants", ("Etudiants", {d: n - 5 for d, n in alld.items()})),
                               ("Inscriptions principales", ("Inscriptions principales", alld)),
                               *[(("UNIVERSITE EXEMPLE NORD (SITE)" if k == "UNIVERSITE EXEMPLE NORD" else k), v) for k, v in unis.items()]]:
        n_all = sum(ds.values())
        tot = [total, "Toutes disciplines", n_all] + ["ns"] * 20 if n_all < 100 else None
        lines = [reo_line(short, d, n) for d, n in ds.items()]
        if tot is None:
            num = [sum(l[i] for l in lines if not isinstance(l[3], str)) for i in range(3, 23, 2)]
            tot = [total, "Toutes disciplines", n_all]
            for x in num:
                tot += [x, round(100 * x / n_all, 1)]
        ws.append(tot)
        for l in lines:
            ws.append(l)
            plan["reo"][(total, l[1])] = l
        ws.append([])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return plan


REO_SYN_ALIASES = {"UNIVERSITE EXEMPLE NORD (SITE)": "UNIVERSITE EXEMPLE NORD"}
PASSERELLES_PAYSAGE = {"UNIVERSITE EXEMPLE PARIS (CITE)": "PARIS", "UNIVERSITE EXEMPLE NORD": "NORD1"}


def make_monmaster_raw(raw_dir: Path) -> dict:
    """monmaster_2024_2025.parquet with the real column names: programmes of the Paris university (droit,
    sciences, a MEEF master filed under "Sciences humaines et sociales", a mixed droit-éco one), one of
    another establishment, two sessions. Origins add up to the admits; candidates split main/complementary."""
    import pandas as pd
    rows = []
    offers = [  # id, ifc, mention, disci_lib, disci_master, can_pp, can_pc, prop, accept, etab, origins (lg3, lp3, but3, master, autre, noninscri)
        ("PARIS", "P-DROIT-1", "DROIT PRIVE", "Droit, sciences politiques", "Droit et sc. politiques", 900, 10, 120, 40, 22, (34, 0, 1, 2, 1, 2)),
        ("PARIS", "P-DROIT-2", "DROIT DES AFFAIRES", "Droit, sciences politiques", "Droit et sc. politiques", 1500, 0, 200, 60, 20, (50, 2, 3, 1, 2, 2)),
        ("PARIS", "P-SCI-1", "CHIMIE", "Sciences fondamentales et applications", "Sc. fondamentales et appliquées (dont STAPS)", 200, 5, 60, 25, 15, (20, 1, 2, 0, 1, 1)),
        ("PARIS", "P-MEEF", "MEEF PREMIER DEGRE", "Sciences humaines et sociales", "MEEF", 400, 0, 90, 50, 10, (40, 0, 0, 5, 2, 3)),
        ("PARIS", "P-MIX", "DROIT ECONOMIE GESTION", "Pluridisciplinaire droit, sciences économiques, AES", "Economie, gestion et AES", 300, 0, 50, 20, 5, (18, 0, 0, 1, 0, 1)),
        ("NORD1", "N-DROIT-1", "DROIT PUBLIC", "Droit, sciences politiques", "Droit et sc. politiques", 600, 20, 90, 30, 12, (25, 1, 0, 1, 1, 2)),
    ]
    for session in ("2024", "2025"):
        for pid, ifc, mention, dl, dm, cpp, cpc, prop, acc, etab, orig in offers:
            shift = 0 if session == "2025" else -2
            o = list(orig)
            o[0] += shift
            r = {"session": session, "ifc": f"{ifc}-{session}", "id_paysage": pid, "id_paysage_actuel": pid, "eta_nom": f"Etablissement {pid}",
                 "inm": mention[:6], "mention": mention, "parcours": "Parcours", "alternance": "0", "disci_lib": dl, "disci_master": dm,
                 "n_can_pp": cpp, "n_can_pc": cpc, "n_prop_total": prop, "n_accept_total": acc + shift, "n_accept_etab": etab}
            for k, v in zip(("lg3", "lp3", "but3", "master", "autre", "noninscri"), o):
                r[f"n_accept_{k}_total"] = v
            rows.append(r)
    raw_dir.mkdir(parents=True, exist_ok=True)
    pd.DataFrame(rows).to_parquet(raw_dir / "monmaster_2024_2025.parquet", index=False)
    return {"offers": offers}


# --------------------------------------------------------------------------- sprint 6 workbooks
POURSUITE_ALIASES = {"UNIV EXEMPLE LYON (EPE)": "UNIVERSITE EXEMPLE LYON", "UNIVERSITE EXEMPLE PARIS": "UNIVERSITE EXEMPLE PARIS (CITE)",
                     "UNIVERSITE EXEMPLE PARIS (EPE)": "UNIVERSITE EXEMPLE PARIS (CITE)"}
POURSUITE_PAYSAGE = {"PARIS": "UNIVERSITE EXEMPLE PARIS (CITE)", "LYON1": "UNIVERSITE EXEMPLE LYON"}
POURSUITE_RULES = {r"^lycee exemple marseille": "UNIVERSITE EXEMPLE NORD"}


def make_poursuite_xlsx(d: Path) -> dict:
    """The five sprint-6 workbooks with the real files' layout and quirks: titles in A1, "Retour au
    sommaire", national "Etudiants" (not additive) and "Inscriptions principales" (= sum) rows, "ns"
    below 100 entrants, (EPE) spellings, discipline rows carrying a shorter name than their block's
    total row, a newline inside a discipline label, trailing spaces in specialty labels, footers."""
    import openpyxl
    d.mkdir(parents=True, exist_ok=True)
    plan = {}

    def wb_with(sheets):
        wb = openpyxl.Workbook()
        wb.active.title = "Sommaire"
        wb.active.append(["Sommaire"])
        for title, rows in sheets:
            ws = wb.create_sheet(title)
            for r in rows:
                ws.append(list(r))
        return wb

    # BUT per university: (row name, entrants, graduates in 3 years, simulated, continuation %, master, commerce, ingé)
    but = [("UNIVERSITE EXEMPLE PARIS (CITE)", 400, 240, 55.0, 70.0, 30.0, 20.0, 10.0),
           ("UNIV EXEMPLE LYON (EPE)", 300, 150, 52.0, 60.0, 20.0, 15.0, 12.0),
           ("UNIVERSITE EXEMPLE NORD", 60, 30, None, None, None, None, None)]          # < 100: "ns"
    n, k = sum(b[1] for b in but), sum(b[2] for b in but)
    head = ["Nom établissement", "Premières\ninscriptions\nen 2021", "Réussite\nen 3 ans", "Réussite\nen 3 ans\nTaux\nobservé",
            "Réussite\nen 3 ans\nTaux\nsimulé", "Réussite\nen 3 ans\nValeur\najoutée\n(en\npoints)"]
    rows = [["Réussite au BUT des bacheliers 2021 inscrits pour la première fois en première année de BUT en 2021, par établissement"],
            ["Champ : Universités, France entière"], ["Source: MESRE-SIES, système d'information SISE"], ["Retour au sommaire"], head,
            ["Etudiants", n - 5, k - 3, 100 * (k - 3) / (n - 5), ".", "."], ["Inscriptions principales", n, k, 100 * k / n, ".", "."]]
    for name, e, g, sim, *_ in but:
        rows.append([name, e, *(["ns"] * 4 if e < 100 else [g, round(100 * g / e, 1), sim, round(round(100 * g / e, 1) - sim, 1)])])
    rows += [[], ["ns : résultat non significatif, en raison d'effectifs trop faibles."]]
    ph = ["Nom établissement", "Effectif", "Répartition des\ndiplômés (%)", "Poursuite d'études\n(%)", "Dont Master\n(%)",
          "Dont école de commerce\n(%)", "Dont école d'ingénieurs \n(%)"]
    prow = [["Poursuites d’études en 2024 des bacheliers 2021 inscrits pour la première fois en première année de BUT en 2021 et diplômés en trois ans, par établissement"],
            ["Champ : Universités, France entière"], ["Source: MESRE-SIES, système d'information SISE"], ["Retour au sommaire"], ph,
            ["Ensemble", k, 100, 65.0, 25.0, 17.0, 11.0]]
    for name, e, g, sim, c, m, co, ing in but:
        prow.append([name, g, 100 * g / k, *(["ns"] * 4 if g < 100 else [c, m, co, ing])])
    p = d / "nf-sies-test-but-etab.xlsx"
    wb_with([("Définitions et méthodologie", [["Définitions et méthodologie"], ["Retour au sommaire"]]),
             ("Effectif BUT 2021", [["Effectifs"], [], [], [], ["Retour au sommaire"], ["Nom établissement", "Total"]]),
             ("Réussite BUT établissements", rows), ("Poursuite diplômés 3ans", prow)]).save(p)
    plan["but_etab"] = p
    plan["but"] = {"n": n, "k": k, "rows": but}

    t1 = [["Tableau 1 - Réussite au BUT en trois ans selon la filière et la mention de baccalauréat des bacheliers 2021 inscrits en première année à la rentrée 2021 (en %)"],
          ["Retour au sommaire"], [], ["Série de baccalauréat", "Part des inscrits", "Taux de passage en 2ème année", "Taux de passage en 3ème année", "Taux de réussite en 3 ans"],
          ["Ensemble bac général*", 57.7, 81.2, 64.7, 62.4], ["dont MATHS / PHY_CHI", 16.1, 79.8, 60.9, 58.7],
          ["dont MATHS / SES", 5.9, 87.9, 67.3, 65.2], ["Ensemble bac technologique", 41.0, 59.1, 47.9, 44.9],
          ["dont STMG (1)", 20.8, 62.8, 50.2, 46.9], ["dont autres bac techno.", 20.2, 55.3, 45.4, 42.9],
          ["Ensemble bac professionnel", 1.3, 53.8, 40.8, 37.7], ["Mention obtenue"], ["Très bien", 4.4, 86.0, 68.0, 66.4],
          ["Bien", 28.1, 83.0, 66.1, 64.2], ["Assez bien", 47.3, 71.5, 57.6, 54.9], ["Passable 1er groupe", 17.9, 53.3, 42.9, 39.6],
          ["Passable 2nd groupe", 0.6, 33.9, 24.1, 18.6], ["Inconnue", 1.7, 67.6, 52.7, 49.7], ["Ensemble", 100, 71.8, 57.5, 100 * k / n],
          ["*Les taux sont présentés lorsque la combinaison de spécialités représente plus de 1 % des inscrits."]]
    t2 = [["Tableau 2 - Réussite au BUT en trois ans selon la spécialité d'inscription des bacheliers 2021 à la rentrée 2021 (en %)"],
          ["Retour au sommaire"], [], ["Spécialité", "Part des inscrits", "Taux de passage en 2ème année", "Taux de passage en 3ème année", "Taux de réussite en 3 ans"],
          ["Secteur de la production", 40.0, 68.5, 55.0, 53.1], ["Génie mécanique et productique ", 6.3, 70.3, 57.2, 54.5],
          ["Statistique et informatique décisionnelle", 1.0, 72.9, 54.8, 53.0], ["Secteur des services", 60.0, 74.0, 59.2, 56.2],
          ["Gestion des entreprises et des administrations ", 17.7, 71.5, 52.9, 51.1], ["Ensemble", 100, 71.8, 57.5, 100 * k / n],
          ["Source : MESRE-SIES, système d'information SISE"]]
    p = d / "nf-sies-test-but-nat.xlsx"
    wb_with([("Définitions et méthodologie", [["x"]]), ("Tableau 1", t1), ("Tableau 2", t2)]).save(p)
    plan["but_nat"] = p

    # master per university x discipline: block total name, discipline-row name, {discipline: (entrants, success, simulated)}
    blocks = [("UNIVERSITE EXEMPLE PARIS (EPE)", "UNIVERSITE EXEMPLE PARIS",
               {"Droit, sc. politiques": (300, 240, 78.0), "Sciences-santé": (200, 150, 80.0), "STAPS": (40, 30, None)}),
              ("UNIVERSITE EXEMPLE NORD", "UNIVERSITE EXEMPLE NORD", {"Droit, sc. politiques": (150, 105, 75.0)})]

    def mtab(title, head, line):
        out = [[title], ["retour au sommaire"], [], head]
        nat = {}
        for _, _, discs in blocks:
            for dname, (e, s, _) in discs.items():
                nat.setdefault(dname, [0, 0])
                nat[dname][0] += e
                nat[dname][1] += s
        te, ts = sum(v[0] for v in nat.values()), sum(v[1] for v in nat.values())
        for natname, shift in (("Etudiants", -3), ("Inscriptions principales", 0)):
            out.append(line("Toutes disciplines", natname, te + shift, ts + shift, None, national=True))
            for dname, (e, s) in nat.items():
                out.append(line(dname, natname, e + shift, s + shift, None, national=True))
            out.append([])
        for total, short, discs in blocks:
            e_, s_ = sum(v[0] for v in discs.values()), sum(v[1] for v in discs.values())
            out.append(line("Toutes disciplines", total, e_, s_, 76.0))
            for dname, (e, s, sim) in discs.items():
                out.append(line(dname, short, e, s, sim))
        out += [[], ["Champ : Universités-établissements assimilés, France entière"]]
        return out

    def l1(disc, name, e, s, sim, national=False):
        if e < 100:
            return [disc, name, e, "ns", "ns", "ns", "ns"]
        obs = round(100 * s / e, 1)
        return [disc, name, e, s, obs, "." if national else sim, "." if national else round(obs - sim, 1)]

    def l3(disc, name, e, s, sim, national=False):
        disc = disc.replace("Droit, sc. politiques", "Droit, Sc.\npolitiques")
        if e < 100:
            return [disc, name, e] + ["ns"] * 16
        obs = round(100 * s / e, 1)
        rest = 100 - obs
        red, reo, sortie = round(rest * 0.5, 1), round(rest * 0.2, 1), None
        sortie = round(100 - obs - red - reo, 1)
        return [disc, name, e, s, obs, "." if national else sim, "." if national else round(obs - sim, 1),
                round(e * red / 100), red, round(e * reo / 100), reo, 0, 0.0, 0, 0.0, 0, 0.0, round(e * sortie / 100), sortie]

    h1 = ["Disciplines en M1", "Nom établissement", "Effectif 2021", "Réussite à 2 ou 3 ans", "Taux observé", "Taux simulé", "Valeur ajoutée (en points)"]
    h3 = ["Discipline", "Etablissement d'inscription en M1", "Premières inscriptions en M1", "Passage\nen M2", "Passage\nen M2\nTaux\nobservé",
          "Passage\nen M2\nTaux\nsimulé", "Passage en M2 Valeur ajoutée (en points)", "Redoublement en M1", "Taux de redoublement en M1 ",
          "Réorientation à l'université", "Taux de réorientation à l'université", "Réorientation en école d'ingénieur ou ENS",
          "Taux de réorientation en école d'ingénieur ou ENS", "Réorientation en école de management", "Taux de réorientation en école de management ",
          "Autres réorientations", "Taux autres réorientations\n", "Sorties", "Taux de sortie "]
    p = d / "nf-sies-test-master-etab.xlsx"
    wb_with([("Méthodologie", [["Méthodologie"]]),
             ("Tableau 1", mtab("Tableau 1 - Réussite en 2 ou 3 ans par discipline et établissement des étudiants inscrits pour la première fois en M1 en 2021", h1, l1)),
             ("Tableau 2", [["Tableau 2 - Réussite en 2 ans par discipline et établissement des étudiants inscrits pour la première fois en M1 en 2022"]]),
             ("Tableau 3", mtab("Tableau 3 - Devenir en 2024-2025 des étudiants inscrits pour la première fois en M1 en 2023", h3, l3))]).save(p)
    plan["master_etab"] = p
    plan["master"] = blocks

    t3 = [["Tableau 3 - Poursuite en master universitaire des diplômés de licence générale (en %)"], ["retour au sommaire"], [],
          ["Discipline de LG", "Poursuite en master des diplômés de LG"], [None, "diplômés 2021", "diplômés 2022", "diplômés 2023", "diplômés 2024"],
          ["Droit, sciences politiques", 68.6, 71.4, 75.2, 79.5], ["Sciences-santé", 67.5, 66.5, 68.9, 67.8],
          ["Sciences et techniques des activités physiques et sportives (STAPS)", 45, 44.6, 49, 50.9], ["Ensemble", 57.8, 57.6, 60.8, 63.3],
          ["Source : MESRE-SIES, système d'information sur le suivi de l'étudiant (SISE)"]]
    p = d / "nf-sies-test-master-nat.xlsx"
    wb_with([("Tableau 1", [["Tableau 1 - Réussite en master universitaire en deux ou trois ans"]]), ("Tableau 3", t3)]).save(p)
    plan["master_nat"] = p

    ch = [["Parcours et réussite en licence et en master universitaire des bacheliers 2017 inscrits en L1 à la rentrée 2017"], ["retour au sommaire"], [],
          ["Caractéristiques                   des étudiants", "Part des inscrits", "Réussite en licence en 3 ou 4 ans",
           "Passage en master après 3 ou 4 ans d'études", "Réussite au master après 5 ou 6 ans d'études", "Réussite au master après 5, 6 ans ou 7 ans d'études"],
          ["Homme", 40.8, 38.4, 24.7, 19.4, 22.7], ["Femme", 59.2, 51.6, 34.8, 28.5, 32.1], ["Type de baccalauréat"],
          ["Général", 80.3, 54.1, 36.5, 29.7, 33.7], ["Technologique", 13.3, 17.6, 8.8, 6.2, 7.9], ["Professionnel", 6.4, 6.4, 3, 1.6, 2.2],
          ["Mention obtenue au baccalauréat"], ["Très bien", 6, 78.9, 62.1, 54.9, 58.2], ["Assez bien", 27, 55.3, 37.1, 30.3, 34.3],
          ["Passable premier groupe", 36.5, 34.7, 20.2, 15.2, 18.5], ["Discipline de L1"], ["Droit, sciences politiques", 18.7, 49.7, 38.5, 33.3, 37.6],
          ["SVT", 5.6, 47.1, 32.6, 28.5, 32.8], ["Ensemble", 100, 46.2, 30.7, 24.8, 28.3], ["Source : MESRE-SIES"]]
    p = d / "nf-sies-test-master-chain.xlsx"
    wb_with([("Méthodologie", [["Méthodologie"]]), ("Tableau 1", [["Poursuite d'études en 2024-2025 des diplômés 2024 de licence générale"]]),
             ("Tableau 2", ch)]).save(p)
    plan["master_chain"] = p
    plan["sources"] = {k: plan[k] for k in ("but_etab", "but_nat", "master_etab", "master_nat", "master_chain")}
    return plan
