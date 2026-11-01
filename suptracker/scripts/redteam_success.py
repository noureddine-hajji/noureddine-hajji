"""Red team of B1 "Réussite en licence": SIES success after L1, per university x broad discipline.

    python scripts/redteam_success.py [--db data/suptracker.duckdb] [--xlsx data/sources/...xlsx] [--real]

The checks re-read the workbook independently (openpyxl, their own block grouping), rebuild the
B1 tables in memory from the warehouse's formations, and compare: identities SIES must satisfy,
stored = published, masked never 0, the national reference, the university links (truth table of
the hard names), the discipline map, the API block, the UI wording and escaping. A mutation test
seeds bugs and requires each to be caught.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

import duckdb

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

from suptracker import success as SU  # noqa: E402

CHECKS = []
APP_JS = ROOT / "suptracker" / "static" / "app.js"
DISC = {"Toutes disciplines": "toutes", "Droit, Sc. politiques": "droit", "Economie, AES": "eco",
        "Arts, lettres, langues, SHS": "alsh", "Sciences-santé": "sciences", "STAPS": "staps"}

# Parcoursup establishment names that are easy to get wrong -> SIES university (None: not linkable)
TRUTH = {
    "Université de Montpellier, Antenne de Perpignan": "UNIVERSITE MONTPELLIER",
    "Université de Perpignan Via Domitia": "UNIVERSITE PERPIGNAN",
    "Université de Montpellier, Antenne de Nîmes": "UNIVERSITE MONTPELLIER",
    "Université de Nîmes": "UNIVERSITE NIMES",
    "Université Paul Valery - Centre Universitaire Du Guesclin - Béziers": "UNIVERSITE MONTPELLIER 3",
    "Université Paris 1 Panthéon Sorbonne": "UNIVERSITE PARIS 1 (PANTHEON SORBONNE)",
    "Sorbonne université": "SORBONNE UNIVERSITE",
    "Observatoire océanologique de Roscoff - Sorbonne Université": "SORBONNE UNIVERSITE",
    "Université Sorbonne Nouvelle Paris 3": "UNIVERSITE PARIS 3 (SORBONNE NOUVELLE)",
    "Université Sorbonne Paris Nord - Antenne de Bobigny": "UNIVERSITE PARIS 13 (SORBONNE PARIS NORD)",
    "Université Paris- Est-Créteil Val de Marne - UPEC (Paris 12)": "UNIVERSITE PARIS-EST CRETEIL",
    "Université Panthéon- Assas Paris2 - Antenne Melun": "UNIVERSITE PARIS 2 (PANTHEON ASSAS)",
    "Université Paris Panthéon Assas": "UNIVERSITE PARIS 2 (PANTHEON ASSAS)",
    "Université Lumière - Lyon 2": "UNIVERSITE LYON 2 (LUMIERE)",
    "Université Lyon 2 - Campus Porte des Alpes": "UNIVERSITE LYON 2 (LUMIERE)",
    "Campus de Bourg-en-Bresse - Université Lyon 3": "UNIVERSITE LYON 3 (JEAN MOULIN)",
    "Université Claude Bernard Lyon 1": "UNIVERSITE LYON 1 (CLAUDE BERNARD)",
    "Université Toulouse III - Université Toulouse": "UNIVERSITE TOULOUSE 3 (PAUL SABATIER)",
    "Université Toulouse 2 - Antenne de Foix": "UNIVERSITE TOULOUSE 2 (JEAN JAURES)",
    "Université Toulouse Capitole - Antenne de Montauban": "UNIVERSITE TOULOUSE 1 (CAPITOLE)",
    "Université de Rennes (EPE) - Site de Saint-Brieuc": "UNIVERSITE DE RENNES (RENNES 1)",
    "UFR STAPS - Université de Rennes 2 - Site de Saint-Brieuc": "UNIVERSITE RENNES 2",
    "Université Bordeaux Montaigne - Antenne d'Agen": "UNIVERSITE BORDEAUX 3",
    "Université de Bordeaux - Antenne d'Agen": "UNIVERSITE DE BORDEAUX",
    "Universite de Lille - Centre Universitaire de Cambrai": "UNIVERSITE LILLE",
    "Université Polytechnique Hauts-de-France - Site de CAMBRAI": "UNIVERSITE POLYTECHNIQUE DES HAUTS DE FRANCE",
    "Université d'Evry Val d'Essonne": "UNIVERSITE EVRY VAL D ESSONNE",
    "Université Paris-Saclay, Campus d'Orsay": "UNIVERSITE PARIS SACLAY",
    "Université Paris-Saclay, Campus d’Evry": None,
    "Ecole Universitaire de premier cycle - Campus d'Evry Université Paris-Saclay": None,
    "Université Paris-Saclay - Campus de Versailles": None,
    "Ecole Universitaire de premier cycle - Campus de Guyancourt, Versailles Saint Quentin en Yvelines, Université Paris-Saclay": None,
    "Université de Versailles - Saint Quentin en Yvelines": "UNIVERSITE VERSAILLES ST QUENTIN",
    "Université Marie et Louis Pasteur - Université de Franche-Comté (Belfort)": "UNIVERSITE DE BESANCON",
    "Université Bourgogne Europe (ex université de Bourgogne)": "UNIVERSITE DIJON",
    "Universite Jean Monnet, Saint-Etienne": "UNIVERSITE SAINT ETIENNE",
    "Institut national universitaire Champollion - Campus de Rodez": "INU CHAMPOLLION",
    "INALCO": None,
    "Institut d'études politiques de Paris - Sciences Po": None,
    "Lycée Henri IV": None,
    "École d'économie et de sciences sociales quantitatives - Toulouse School of Economics (TSE)": None,
    "Université Paris Cité": "UNIVERSITE PARIS CITE",
    "Université Paris 8": "UNIVERSITE PARIS 8 (VINCENNES SAINT-DENIS)",
    "Université de Mayotte": "CUFR MAYOTTE",
}
PROGRAMME_TRUTH = {
    "Psychologie": "alsh", "Science politique": "droit", "Gestion": "eco", "Administration économique et sociale": "eco",
    "Sciences et Techniques des Activités Physiques et Sportives": "staps", "Sciences pour la santé": "sciences",
    "Géographie et aménagement": "alsh", "Information et communication": "alsh", "Droit": "droit",
    "Mathématiques et informatique appliquées aux sciences humaines et sociales": None, "Droit, Economie": None,
    "Sciences sanitaires et sociales": None, "Economie et gestion": "eco", "Informatique": "sciences",
}


def check(name):
    def deco(fn):
        fn.check_name = name
        CHECKS.append(fn)
        return fn
    return deco


class Ctx:
    """The workbook re-read independently, and B1 tables rebuilt in memory."""

    def __init__(self, db: Path, xlsx: Path, table: dict | None, real: bool):
        self.db, self.xlsx, self.table, self.real = db, xlsx, table, real
        self.raw = read_raw(xlsx)
        self.rebuild()

    def rebuild(self):
        src = duckdb.connect(str(self.db), read_only=True)
        latest = src.execute("SELECT max(session) FROM core.formation_year").fetchone()[0]
        forms = src.execute("""SELECT session, cod_aff_form, g_ea_lib_vx, fil_lib_voe_acc, fili, contrat_etab
                               FROM core.formation_year WHERE session = ?""", [latest]).fetchall()
        self.has_stored = bool(src.execute("""SELECT count(*) FROM information_schema.tables
            WHERE table_schema = 'core' AND table_name = 'licence_success_link'""").fetchone()[0])
        self.stored_links = (src.execute("SELECT cod_aff_form, etab, discipline FROM core.licence_success_link").fetchall()
                             if self.has_stored else None)
        src.close()
        con = duckdb.connect()
        con.execute("CREATE SCHEMA core")
        con.execute("""CREATE TABLE core.formation_year (session INTEGER, cod_aff_form VARCHAR, g_ea_lib_vx VARCHAR,
                       fil_lib_voe_acc VARCHAR, fili VARCHAR, contrat_etab VARCHAR)""")
        con.executemany("INSERT INTO core.formation_year VALUES (?, ?, ?, ?, ?, ?)", forms)
        self.n_linked = SU.attach(con, self.xlsx, self.table)
        self.con, self.latest, self.forms = con, latest, forms


def _cell(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    return str(v).strip() if v is not None else None


def read_raw(xlsx: Path) -> dict:
    """{sheet kind: {(block, discipline): {header: value}}} with its own grouping logic."""
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    kinds = {"devenir": "passage", "réussite en licence en 3 ou 4 ans": "reussite34", "réussite en licence en 3 ans": "reussite3"}
    out = {}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        title = str(rows[0][0] or "").lower() if rows and rows[0] else ""
        kind = next((k for t, k in kinds.items() if title.startswith(t)), None)
        if not kind:
            continue
        head_i = next(i for i, r in enumerate(rows) if r and str(r[0] or "").startswith("Nom établissement"))
        head = [re.sub(r"\s+", " ", str(h or "")).strip() for h in rows[head_i]]
        cur, blocks = None, {}
        for r in rows[head_i + 1:]:
            if not r or not r[0] or not r[1]:
                continue
            d = DISC.get(re.sub(r"\s+", " ", str(r[1])).strip())
            if d is None:
                continue
            if d == "toutes":
                cur = SU.SIES_ALIASES.get(str(r[0]).strip(), str(r[0]).strip())
            blocks[(cur, d)] = {h: _cell(v) for h, v in zip(head, r) if h}
        out[kind] = blocks
    return out


def _f(x):
    return x if isinstance(x, float) else None


RAW_KEYS = {  # stored column -> raw header, per kind
    "passage": {"inscrits": "Inscrits en L1", "taux": "Taux de passage en L2", "taux_simule": "Taux de passage en L2 simulé",
                "va": "Passage en L2 - Valeur Ajoutée (en points)", "taux_redoublement": "Taux de redoublement en L1",
                "taux_reorientation": "Taux de réorientation", "taux_sortie": "Taux de sortie de l'enseignement supérieur",
                "n_passage": "Passage en L2", "n_redoublement": "Redoublement en L1", "n_reorientation": "Réorientation",
                "n_sortie": "Sortie de l'enseignement supérieur"},
    "reussite3": {"inscrits": "Inscrits en L1", "taux": "Taux de réussite à 3 ans", "taux_simule": "Taux simulé",
                  "va": "Valeur ajoutée (en points)", "n_reussite": "Réussite en 3 ans"},
    "reussite34": {"inscrits": "Inscrits en L1", "taux": "Réussite en 3 ou 4 ans (taux observé)",
                   "taux_simule": "Réussite en 3 ou 4 ans (taux simulé)", "va": "Valeur ajoutée réussite en 3 ou 4 ans (en points)",
                   "taux3": "Réussite en 3 ans (taux observé)", "taux4": "Réussite en 4 ans (taux observé)",
                   "n_reussite": "Réussite en 3 ou 4 ans"},
}


@check("S1 SIES identities hold in the workbook (sums, partition, value added, rates)")
def s1(c: Ctx):
    bad = []
    for kind, blocks in c.raw.items():
        by = {}
        for (b, d), v in blocks.items():
            by.setdefault(b, {})[d] = v
        unis = [b for b in by if b not in ("Etudiants", SU.NATIONAL_ROW)]
        for b, ds in by.items():
            parts = sum(_f(v.get("Inscrits en L1")) or 0 for d, v in ds.items() if d != "toutes")
            if abs(parts - (_f(ds["toutes"].get("Inscrits en L1")) or 0)) > 0.5:
                bad.append((kind, b, "discipline sum"))
        for d in DISC.values():
            tot = sum(_f(by[b].get(d, {}).get("Inscrits en L1")) or 0 for b in unis)
            nat = _f(by[SU.NATIONAL_ROW].get(d, {}).get("Inscrits en L1"))
            if nat is not None and abs(tot - nat) > 0.5:
                bad.append((kind, d, "universities != Inscriptions principales"))
        rk = RAW_KEYS[kind]
        for (b, d), v in blocks.items():
            t, s, va = _f(v.get(rk["taux"])), _f(v.get(rk["taux_simule"])), _f(v.get(rk["va"]))
            if None not in (t, s, va) and abs(t - s - va) > 0.11:
                bad.append((kind, b, d, "va != observed - simulated"))
            n, ins = _f(v.get(rk.get("n_passage") or rk["n_reussite"])), _f(v.get("Inscrits en L1"))
            if None not in (n, ins, t) and ins > 0 and abs(100 * n / ins - t) > 0.06:
                bad.append((kind, b, d, "rate != count / entrants"))
            if kind == "passage" and b not in ("Etudiants", SU.NATIONAL_ROW):
                parts = [_f(v.get(rk[k])) for k in ("n_passage", "n_redoublement", "n_reorientation", "n_sortie")]
                if None not in parts and ins is not None and abs(sum(parts) - ins) > 2.5:
                    bad.append((kind, b, d, f"partition off by {sum(parts) - ins:+.0f}"))
    return not bad, f"{len(bad)} violations {bad[:3]}" if bad else f"{sum(len(b) for b in c.raw.values())} rows consistent"


@check("S2 stored = published, cell by cell; a masked cell is NULL, never 0")
def s2(c: Ctx):
    rows = c.con.execute("SELECT * FROM core.licence_success").fetchall()
    cols = [x[0] for x in c.con.description]
    wrong, n, masked_ok, masked = [], 0, 0, 0
    stored = set()
    for r in rows:
        r = dict(zip(cols, r))
        key = (SU.NATIONAL_ROW if r["etab"] == SU.NATIONAL else r["etab"], r["discipline"])
        stored.add((r["kind"], key))
        raw = c.raw[r["kind"]].get(key)
        if raw is None:
            wrong.append((r["kind"], key, "not in workbook"))
            continue
        for col, h in RAW_KEYS[r["kind"]].items():
            n += 1
            rv = raw.get(h)
            if isinstance(rv, float):
                if r[col] is None or abs(r[col] - rv) > 1e-9:
                    wrong.append((r["kind"], key, col, rv, r[col]))
            else:
                masked += 1
                masked_ok += r[col] is None
    missing = [(k, key) for k, blocks in c.raw.items() for key in blocks
               if key[0] != "Etudiants" and (k, key) not in stored]
    ok = not wrong and masked_ok == masked and not missing
    return ok, (f"{n} cells compared, {len(wrong)} differ {wrong[:2]}; masked kept NULL {masked_ok}/{masked}; "
                f"rows not stored {missing[:2]}")


@check("S3 the national reference is 'Inscriptions principales' (what the universities add up to), never 'Etudiants'")
def s3(c: Ctx):
    nat = dict(c.con.execute("""SELECT discipline, inscrits FROM core.licence_success WHERE kind = 'passage' AND etab = ?""",
                             [SU.NATIONAL]).fetchall())
    uni = dict(c.con.execute("""SELECT discipline, sum(inscrits) FROM core.licence_success WHERE kind = 'passage'
                                AND etab <> ? GROUP BY 1""", [SU.NATIONAL]).fetchall())
    etu = c.con.execute("SELECT count(*) FROM core.licence_success WHERE etab_row = 'Etudiants'").fetchone()[0]
    bad = [d for d in uni if d != "toutes" and abs((nat.get(d) or 0) - uni[d]) > 0.5]
    return not bad and not etu and nat, f"disciplines off {bad}; 'Etudiants' rows stored: {etu}"


@check("S4 every SIES university has a link rule, a display name and all three cohorts")
def s4(c: Ctx):
    table = c.table or SU.UNIVERSITIES
    blocks = {b for k in c.raw.values() for (b, _) in k} - {"Etudiants", SU.NATIONAL_ROW}
    no_rule = sorted(blocks - set(table))
    no_label = sorted(set(SU.UNIVERSITIES) - set(SU.UNIVERSITY_LABEL)) if c.table is None else []
    partial = sorted(b for b in blocks if sum(1 for k in c.raw.values() if (b, "toutes") in k) != len(c.raw))
    stale = sorted(set(table) - blocks)
    return not (no_rule or no_label or partial or stale), f"no rule {no_rule}; no label {no_label}; missing cohorts {partial}; rules for absent universities {stale}"


@check("S5 university links: hard names resolve as expected, a name never links two universities, only public licences")
def s5(c: Ctx):
    bad = []
    if c.real:
        for name, want in TRUTH.items():
            got, hits = SU.university_of(name, c.table)
            if got != want:
                bad.append((name, want, got, hits))
    links = c.con.execute("SELECT l.cod_aff_form, l.etab, l.matches FROM core.licence_success_link l").fetchall()
    multi = [x for x in links if x[1] is not None and "|" in (x[2] or "")]
    priv = c.con.execute("""SELECT count(*) FROM core.licence_success_link l JOIN core.formation_year f USING (cod_aff_form)
                            WHERE f.contrat_etab <> 'Public' OR f.fili NOT IN ('Licence', 'Licence_Las')""").fetchone()[0]
    n_pub = sum(1 for f in c.forms if f[5] == "Public" and f[4] in SU.LICENCE_FILI)
    cover = 100 * c.n_linked / n_pub if n_pub else 0
    ok = not bad and not multi and not priv and (cover >= 90 if c.real else c.n_linked > 0)
    return ok, f"truth-table misses {bad[:2]}; links with 2 universities {len(multi)}; non-public or non-licence {priv}; coverage {cover:.1f} % of {n_pub} public licences"


@check("S6 discipline map: closed set, expected verdicts, double or cross-field licences left unclassified")
def s6(c: Ctx):
    wrong = [(p, w, SU.discipline_of(p)) for p, w in PROGRAMME_TRUTH.items() if SU.discipline_of(p) != w]
    dup = [p for p in SU.PROGRAMME_DISCIPLINE if sum(p in ps for ps in SU.DISCIPLINE_OF_PROGRAMME.values()) > 1]
    outside = set(SU.DISCIPLINE_OF_PROGRAMME) - (set(SU.DISCIPLINES) - {"toutes"})
    crossed = [p for p in SU.PROGRAMME_DISCIPLINE if re.search(r",|/", p) and
               len({SU.discipline_of(x.strip()) for x in re.split(r",|/", p) if SU.discipline_of(x.strip())}) > 1]
    return not (wrong or dup or outside or crossed), f"wrong {wrong}; listed twice {dup}; codes outside SIES {outside}; cross-field classified {crossed}"


@check("S7 the formation block shows the university-discipline figures and the same-discipline national ones")
def s7(c: Ctx):
    rows = c.con.execute("""SELECT l.cod_aff_form, l.etab, l.discipline, f.fili, f.contrat_etab FROM core.licence_success_link l
        JOIN core.formation_year f USING (cod_aff_form) ORDER BY hash(l.cod_aff_form) LIMIT 300""").fetchall()
    bad, n = [], 0
    for cod, etab, disc, fili, contrat in rows:
        b = SU.for_formation(c.con, cod, fili, contrat)
        n += 1
        if etab is None:
            bad += [] if b["status"] == "non_rattachee" and not b["kinds"] else [(cod, "unlinked shown")]
            continue
        if disc is None:
            bad += [] if b["status"] == "discipline_non_classee" and not b["kinds"] else [(cod, "unclassified shown")]
            continue
        for kind, u in b["kinds"].items():
            if (u["etab"], u["discipline"]) != (etab, disc):
                bad.append((cod, kind, "other university or discipline"))
        for kind, nrow in b["national"].items():
            if nrow["etab"] != SU.NATIONAL or nrow["discipline"] != disc:
                bad.append((cod, kind, "national of another discipline"))
    other = c.con.execute("SELECT cod_aff_form, fili, contrat_etab FROM core.formation_year WHERE fili NOT IN ('Licence', 'Licence_Las') LIMIT 50").fetchall()
    leak = [x[0] for x in other if SU.for_formation(c.con, x[0], x[1], x[2]) is not None]
    priv = c.con.execute("SELECT cod_aff_form, fili, contrat_etab FROM core.formation_year WHERE fili = 'Licence' AND contrat_etab <> 'Public' LIMIT 20").fetchall()
    priv_bad = [x[0] for x in priv if (SU.for_formation(c.con, *x) or {}).get("status") != "hors_champ"]
    return not bad and not leak and not priv_bad, f"{n} licences checked, {len(bad)} wrong {bad[:2]}; non-licences with a block {leak[:3]}; private not 'hors champ' {priv_bad[:3]}"


def _card(js: str) -> str:
    i = js.index("// ---- B1: success in licence")
    return js[i: js.index("function offerCard", i)]


def ui_card(js: str) -> tuple[bool, str]:
    blk = _card(js)
    # a stored string interpolated as is: ${s.x}, ${s.x || ...} (a ternary that escapes is fine)
    raw = re.findall(r"\$\{s\.(?:etab_label|discipline_label|reason|etab)\s*(?:\}|\|\|)", blk)
    scope = "pas seulement de cette licence" in blk and "tous les néo-bacheliers inscrits en L1" in blk
    nochance = "pas des chances de réussite" in blk and not re.search(r"chances? de r[ée]ussir|vous r[ée]ussirez|garanti", blk, re.I)
    noclass = "ne classe pas les universités" in blk
    va_ctx = re.search(r"u\.taux_simule != null && u\.va != null", blk) is not None
    masked = "SU_MASK" in blk and "petit effectif" in blk
    wired = "${successCard(d.success)}" in js
    ok = not raw and scope and nochance and noclass and va_ctx and masked and wired
    return ok, (f"unescaped {raw[:2]}; scope said: {scope}; no promise: {nochance}; 'ne classe pas': {noclass}; "
                f"value added only beside the expected rate: {va_ctx}; masked wording: {masked}; on the fiche: {wired}")


@check("S8 UI: escaped, says it is the university-discipline cohort, no promise, value added with its expected rate")
def s8(c: Ctx):
    return ui_card(APP_JS.read_text())


@check("S9 the methodology text says what the figures are and are not")
def s9(c: Ctx):
    from suptracker.api import METHODOLOGY
    t = METHODOLOGY.get("success", "")
    need = ["pas sur cette seule licence", "ne classe pas", "pas des chances", "même discipline", "SIES", "petits effectifs"]
    miss = [x for x in need if x not in t]
    return not miss, f"missing {miss}"


@check("S10 the published warehouse carries the same links as a fresh rebuild (when it has the B1 tables)")
def s10(c: Ctx):
    if not c.has_stored:
        return True, "warehouse built before B1: nothing to compare (the API then shows no B1 block)"
    fresh = set(c.con.execute("SELECT cod_aff_form, etab, discipline FROM core.licence_success_link").fetchall())
    stored = set(c.stored_links)
    if not stored and c.table is not None:
        return True, "synthetic warehouse built without the synthetic table"
    diff = fresh ^ stored
    return not diff, f"{len(diff)} links differ {sorted(diff, key=str)[:3]}"


# ------------------------------------------------------------------ mutation test
UI_MUTANTS = [
    ("unescaped university name", "<strong>${esc(s.etab_label)}</strong>", "<strong>${s.etab_label}</strong>"),
    ("scope sentence dropped", "pas seulement de cette licence", "de cette licence"),
    ("value added shown without the expected rate", "u && u.taux_simule != null && u.va != null", "u && u.va != null"),
    ("promise wording", "pas des chances de réussite", "vos chances de réussir"),
    ("card not on the fiche", "${successCard(d.success)}", ""),
]


@check("S11 mutation test: every seeded data, link, discipline and UI bug is caught")
def s11(c: Ctx):
    missed = []
    js = APP_JS.read_text()
    for name, old, new in UI_MUTANTS:
        if old not in js or ui_card(js.replace(old, new, 1))[0]:
            missed.append(name)
    saved = {k: getattr(SU, k) for k in ("_num", "NATIONAL_ROW", "PROGRAMME_DISCIPLINE", "UNIVERSITIES", "SIES_ALIASES", "DISC_CODE", "fold", "parse_workbook")}
    t0 = dict(c.table) if c.table else None

    def ns_zero(v):
        x, why = saved["_num"](v)
        return (0.0, None) if why in ("ns", "nd") else (x, why)

    data_mutants = [
        ("masked 'ns' read as 0", lambda: setattr(SU, "_num", ns_zero), [s2]),
        ("national = 'Etudiants'", lambda: (setattr(SU, "NATIONAL_ROW", "Etudiants"), setattr(SU, "SKIPPED_ROWS", ("Inscriptions principales",))), [s3, s2]),
        ("discipline rows grouped by their own (shorter) name", lambda: setattr(SU, "parse_workbook", lambda p: [
            {**r, "etab": r["etab_row"]} if r["etab"] != SU.NATIONAL else r for r in saved["parse_workbook"](p)]), [s2]),
        ("Psychologie filed under droit", lambda: setattr(SU, "PROGRAMME_DISCIPLINE", {**saved["PROGRAMME_DISCIPLINE"], "Psychologie": "droit"}), [s6]),
        ("double licence classified", lambda: setattr(SU, "PROGRAMME_DISCIPLINE", {**saved["PROGRAMME_DISCIPLINE"], "Droit, Economie": "droit"}), [s6]),
    ]
    if c.real:
        data_mutants += [
            ("other-sheet names of a university not aliased", lambda: setattr(SU, "SIES_ALIASES", {}), [s4, s2]),
            ("Perpignan rule unanchored", lambda: setattr(SU, "UNIVERSITIES", {**saved["UNIVERSITIES"], "UNIVERSITE PERPIGNAN": r"perpignan"}), [s5]),
            ("Paris-Saclay takes its Évry campus", lambda: setattr(SU, "UNIVERSITIES", {**saved["UNIVERSITIES"], "UNIVERSITE PARIS SACLAY": r"paris-saclay"}), [s5]),
            ("typographic apostrophe dropped before folding", lambda: setattr(SU, "fold", lambda x: re.sub(r"\s+", " ", __import__("unicodedata").normalize("NFKD", x or "").encode("ascii", "ignore").decode().lower().replace("’", "'")).strip()), [s5]),
            ("Paris 1 rule catches Paris 12/13", lambda: setattr(SU, "UNIVERSITIES", {**saved["UNIVERSITIES"], "UNIVERSITE PARIS 1 (PANTHEON SORBONNE)": r"paris 1|pantheon sorbonne"}), [s5]),
        ]
    for name, apply, checks in data_mutants:
        try:
            apply()
            SU.DISC_CODE = saved["DISC_CODE"]
            c.raw = read_raw(c.xlsx)
            c.rebuild()
            caught = any(not fn(c)[0] for fn in checks)
        except Exception:  # noqa: BLE001 - a crash is a detection too
            caught = True
        finally:
            for k, v in saved.items():
                setattr(SU, k, v)
            SU.SKIPPED_ROWS = ("Etudiants",)
            c.table = t0
        if not caught:
            missed.append(name)
    c.raw = read_raw(c.xlsx)
    c.rebuild()
    n = len(UI_MUTANTS) + len(data_mutants)
    return not missed, f"{n - len(missed)}/{n} mutants caught" + (f"; missed {missed}" if missed else "")


def run(db: Path, xlsx: Path, table: dict | None = None, real: bool = False) -> list[dict]:
    c = Ctx(db, xlsx, table, real)
    out = []
    for fn in CHECKS:
        try:
            ok, detail = fn(c)
        except Exception as exc:  # noqa: BLE001 - a crash is a failed check
            ok, detail = False, f"crash: {type(exc).__name__}: {exc}"
        out.append({"check": fn.check_name, "ok": bool(ok), "detail": detail})
    return out


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--db", type=Path, default=ROOT / "data" / "suptracker.duckdb")
    ap.add_argument("--xlsx", type=Path, default=None)
    ap.add_argument("--real", action="store_true")
    a = ap.parse_args(argv)
    xlsx = a.xlsx or SU.source_path(ROOT / "data" / "sources")
    res = run(a.db, xlsx, None, a.real)
    for r in res:
        print(("PASS " if r["ok"] else "FAIL ") + f"{r['check']}\n       {r['detail']}")
    bad = [r for r in res if not r["ok"]]
    print(f"\n{len(res) - len(bad)}/{len(res)} checks passed")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
